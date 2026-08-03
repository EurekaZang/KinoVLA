#!/usr/bin/env python3
"""Build a one-slide Chinese graphical summary of the KiNO work.

The slide reuses raster simulator captures and measured sensor charts from the
full Chinese deck.  No SVG artwork is generated.
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts import build_kino_work_chinese_presentation as full


OUT_DIR = ROOT / "paper"
ASSET_DIR = OUT_DIR / "kino_work_onepage_assets"
OUT_PPTX = OUT_DIR / "KiNO_work_onepage_summary_cn.pptx"


def make_task_crops() -> tuple[Path, Path]:
    """Extract only the simulator-pair regions from the paper conflict figure."""

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    source = Image.open(full.FIG / "fig1_bidirectional_conflict.png").convert(
        "RGB"
    )
    width, height = source.size
    split = width // 2
    left = source.crop(
        (
            80,
            round(height * 0.105),
            split - 35,
            round(height * 0.535),
        )
    )
    right = source.crop(
        (
            split + 30,
            round(height * 0.105),
            width - 65,
            round(height * 0.535),
        )
    )
    t2 = ASSET_DIR / "task_t2_scene_pair.png"
    t3 = ASSET_DIR / "task_t3_scene_pair.png"
    left.save(t2, optimize=True)
    right.save(t3, optimize=True)
    return t2, t3


def add_chevron(
    slide,
    x: float,
    y: float,
    w: float = 0.30,
    h: float = 0.42,
    *,
    color: str = full.TEAL,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = full.rgb(color)
    shape.line.fill.background()
    return shape


def add_small_image(
    slide,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    accent: str,
    *,
    focus_y: float = 0.5,
):
    full.add_box(
        slide,
        x,
        y,
        w,
        h,
        fill=full.WHITE,
        line=full.tint(accent, 0.72),
        width=1.1,
        rounded=True,
    )
    full.add_picture_crop(
        slide,
        path,
        x + 0.04,
        y + 0.04,
        w - 0.08,
        h - 0.08,
        focus_y=focus_y,
    )
    full.add_box(
        slide,
        x + 0.10,
        y + 0.10,
        w - 0.20,
        0.29,
        fill=accent,
        line=accent,
        rounded=True,
        transparency=5,
    )
    full.add_text(
        slide,
        label,
        x + 0.14,
        y + 0.105,
        w - 0.28,
        0.27,
        size=9.3,
        color=full.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )


def add_expert_box(
    slide,
    title: str,
    subtitle: str,
    x: float,
    y: float,
    accent: str,
    w: float = 1.24,
):
    full.add_box(
        slide,
        x,
        y,
        w,
        0.56,
        fill=full.tint(accent, 0.90),
        line=full.tint(accent, 0.55),
        width=1.0,
        rounded=True,
    )
    full.add_text(
        slide,
        title,
        x + 0.09,
        y + 0.06,
        w - 0.18,
        0.22,
        size=11.2,
        color=accent,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    full.add_text(
        slide,
        subtitle,
        x + 0.08,
        y + 0.31,
        w - 0.16,
        0.16,
        size=7.4,
        color=full.MUTED,
        align=PP_ALIGN.CENTER,
        margin=0,
    )


def build_slide() -> Presentation:
    task_t2, task_t3 = make_task_crops()
    body_chart, _ = full.make_sensor_chart()

    prs = Presentation()
    prs.slide_width = Inches(full.SW)
    prs.slide_height = Inches(full.SH)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = full.rgb(full.WHITE)

    # Header.
    full.add_badge(
        slide,
        "KINO · ONE-PAGE SUMMARY",
        0.54,
        0.23,
        1.95,
        fill=full.TEAL_LIGHT,
        color=full.TEAL,
        size=8.2,
    )
    full.add_text(
        slide,
        "Feel It, See It, Recover",
        0.54,
        0.61,
        5.30,
        0.46,
        size=29,
        color=full.INK,
        bold=True,
        font=full.FONT_EN,
        margin=0,
    )
    full.add_runs(
        slide,
        [
            ("四足机器人跨模态失败归因：", {"bold": True, "color": full.INK}),
            ("先判断为什么失败", {"bold": True, "color": full.TEAL}),
            ("，再选择如何恢复", {"bold": True, "color": full.INK}),
        ],
        5.70,
        0.72,
        7.05,
        0.34,
        size=16.4,
        valign=MSO_ANCHOR.MIDDLE,
    )
    full.add_line(
        slide,
        0.54,
        1.13,
        12.79,
        1.13,
        color=full.LINE,
        width=0.9,
    )

    # Task panel.
    full.add_box(
        slide,
        0.54,
        1.31,
        4.20,
        3.15,
        fill=full.LIGHT,
        line=full.LINE,
        width=1.0,
        rounded=True,
    )
    full.add_badge(
        slide,
        "01 · TASK",
        0.72,
        1.48,
        0.92,
        fill=full.BLUE_LIGHT,
        color=full.BLUE,
        size=8.6,
    )
    full.add_text(
        slide,
        "同样是“走不动”，原因与恢复动作可能完全相反",
        1.78,
        1.49,
        2.72,
        0.35,
        size=14.2,
        color=full.INK,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )
    full.add_picture_crop(
        slide,
        task_t2,
        0.76,
        1.98,
        3.76,
        0.89,
        focus_y=0.52,
        border=full.BLUE,
        border_width=1.0,
    )
    full.add_badge(
        slide,
        "接触前 T2",
        0.86,
        2.05,
        0.84,
        fill=full.BLUE,
        color=full.WHITE,
        size=8.3,
    )
    full.add_text(
        slide,
        "身体证据相同  →  必须看场景",
        0.82,
        2.91,
        3.64,
        0.25,
        size=12.1,
        color=full.BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    full.add_picture_crop(
        slide,
        task_t3,
        0.76,
        3.26,
        3.76,
        0.72,
        focus_y=0.40,
        border=full.RED,
        border_width=1.0,
    )
    full.add_badge(
        slide,
        "接触后 T3",
        0.86,
        3.33,
        0.84,
        fill=full.RED,
        color=full.WHITE,
        size=8.3,
    )
    full.add_text(
        slide,
        "画面证据相同  →  必须听身体",
        0.82,
        4.04,
        3.64,
        0.25,
        size=12.1,
        color=full.RED,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )

    # Method panel.
    full.add_box(
        slide,
        4.92,
        1.31,
        7.87,
        3.15,
        fill=full.WHITE,
        line=full.LINE,
        width=1.0,
        rounded=True,
    )
    full.add_badge(
        slide,
        "02 · METHOD",
        5.10,
        1.48,
        1.10,
        fill=full.TEAL_LIGHT,
        color=full.TEAL,
        size=8.6,
    )
    full.add_text(
        slide,
        "KiNO：不固定偏向某个传感器，而是学习“这一次该信谁”",
        6.34,
        1.49,
        6.14,
        0.35,
        size=14.2,
        color=full.INK,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )

    add_small_image(
        slide,
        full.operator_path("O2", "front"),
        5.16,
        1.98,
        1.47,
        1.16,
        "5 帧前视 RTX RGB",
        full.BLUE,
        focus_y=0.52,
    )
    add_small_image(
        slide,
        body_chart,
        6.78,
        1.98,
        1.47,
        1.16,
        "19 路身体时间序列",
        full.RED,
        focus_y=0.50,
    )
    full.add_text(
        slide,
        "视觉输入",
        5.16,
        3.22,
        1.47,
        0.23,
        size=10.1,
        color=full.BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    full.add_text(
        slide,
        "本体感受",
        6.78,
        3.22,
        1.47,
        0.23,
        size=10.1,
        color=full.RED,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    add_chevron(slide, 8.34, 2.40, color=full.GRAY)

    add_expert_box(
        slide, "视觉专家", "看环境语义", 8.65, 1.91, full.BLUE, 1.12
    )
    add_expert_box(
        slide, "身体专家", "听接触反应", 8.65, 2.50, full.RED, 1.12
    )
    add_expert_box(
        slide, "联合专家", "整合两路证据", 8.65, 3.09, full.TEAL, 1.12
    )
    add_chevron(slide, 9.84, 2.40, 0.23, color=full.TEAL)

    full.add_box(
        slide,
        10.13,
        1.98,
        1.15,
        1.16,
        fill=full.TEAL_LIGHT,
        line=full.TEAL,
        width=1.4,
        rounded=True,
    )
    full.add_text(
        slide,
        "冲突路由器",
        10.23,
        2.18,
        0.95,
        0.28,
        size=12.2,
        color=full.TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    full.add_text(
        slide,
        "比较置信度\n分歧与不确定性",
        10.22,
        2.55,
        0.97,
        0.43,
        size=8.4,
        color=full.INK,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
        line_spacing=1.0,
    )
    add_chevron(slide, 11.35, 2.40, 0.23, color=full.TEAL)
    add_small_image(
        slide,
        full.DEMO / "adhesion_peel/review_t04.8s.png",
        11.67,
        1.98,
        0.91,
        1.16,
        "安全恢复",
        full.TEAL,
        focus_y=0.50,
    )
    full.add_text(
        slide,
        "原因 → 动作",
        11.67,
        3.22,
        0.91,
        0.23,
        size=9.2,
        color=full.TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    full.add_box(
        slide,
        5.15,
        3.76,
        7.42,
        0.48,
        fill=full.TEAL_LIGHT,
        line=full.TEAL_LIGHT,
        rounded=True,
    )
    full.add_runs(
        slide,
        [
            ("输出：", {"color": full.MUTED}),
            ("异常原因", {"bold": True, "color": full.TEAL}),
            ("  →  ", {"bold": True, "color": full.INK}),
            ("与原因匹配的安全恢复", {"bold": True, "color": full.INK}),
        ],
        5.42,
        3.86,
        6.88,
        0.26,
        size=12.2,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    # Dataset strip.
    full.add_badge(
        slide,
        "03 · DATASET",
        0.54,
        4.68,
        1.08,
        fill=full.PINK_LIGHT,
        color=full.PINK,
        size=8.6,
    )
    full.add_text(
        slide,
        "Kino-Fail：11 个异常算子，覆盖生活 · 生产 · 野外",
        1.76,
        4.68,
        6.00,
        0.34,
        size=15.0,
        color=full.INK,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )
    full.add_text(
        slide,
        "反事实配对  ·  RTX RGB  ·  19 路本体感受  ·  物理/外观独立随机化",
        7.62,
        4.72,
        5.17,
        0.27,
        size=10.2,
        color=full.MUTED,
        bold=True,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )

    short_names = {
        "O1": "低摩擦",
        "O2": "软地面",
        "O3": "地面塌陷",
        "O4": "脚底粘附",
        "O5": "负载变化",
        "O6": "侧向外力",
        "O7": "视物错配",
        "O8": "透明障碍",
        "O9": "底盘托底",
        "O10": "驱力衰减",
        "O11": "本体偏差",
    }
    start_x = 0.54
    gap = 0.065
    width = (12.25 - gap * 10) / 11
    for index in range(1, 12):
        op = f"O{index}"
        x = start_x + (index - 1) * (width + gap)
        full.add_box(
            slide,
            x,
            5.13,
            width,
            1.34,
            fill=full.WHITE,
            line=full.LINE,
            width=0.75,
            rounded=True,
        )
        full.add_picture_crop(
            slide,
            full.operator_path(op, "overview"),
            x + 0.035,
            5.165,
            width - 0.07,
            0.90,
            focus_y=0.52,
        )
        full.add_badge(
            slide,
            op,
            x + 0.08,
            5.23,
            0.35 if index < 10 else 0.42,
            fill=full.BLACK,
            color=full.WHITE,
            size=6.7,
        )
        full.add_text(
            slide,
            short_names[op],
            x + 0.04,
            6.13,
            width - 0.08,
            0.24,
            size=8.3,
            color=full.INK,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )

    full.add_box(
        slide,
        0.54,
        6.68,
        12.25,
        0.47,
        fill=full.LIGHT,
        line=full.LINE,
        width=0.8,
        rounded=True,
    )
    full.add_runs(
        slide,
        [
            ("每个案例：", {"bold": True, "color": full.TEAL}),
            (
                "同场景 / 同起点 / 同路线 / 同相机，",
                {"color": full.INK},
            ),
            (
                "只改变一个物理原因",
                {"bold": True, "color": full.RED},
            ),
            (
                "  →  把“失败归因”从观察性标签变成可验证的物理问题",
                {"color": full.INK},
            ),
        ],
        0.78,
        6.79,
        11.77,
        0.25,
        size=11.4,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    # Minimal footer.
    full.add_line(
        slide,
        0.54,
        7.31,
        12.79,
        7.31,
        color=full.LINE,
        width=0.6,
    )
    full.add_text(
        slide,
        "Feel It. See It. Recover.",
        0.56,
        7.34,
        2.4,
        0.12,
        size=6.8,
        color=full.GRAY,
        font=full.FONT_EN,
        margin=0,
    )
    full.add_text(
        slide,
        "Cross-Modal Failure Attribution for Safe Quadrupedal Navigation Recovery",
        7.30,
        7.34,
        5.47,
        0.12,
        size=6.8,
        color=full.GRAY,
        font=full.FONT_EN,
        align=PP_ALIGN.RIGHT,
        margin=0,
    )
    return prs


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    full.ASSET_DIR.mkdir(parents=True, exist_ok=True)
    full.set_mpl_style()
    deck = build_slide()
    deck.core_properties.title = "KiNO one-page work summary"
    deck.core_properties.subject = "Task, method, and Kino-Fail dataset"
    deck.core_properties.author = "KiNO project"
    deck.save(OUT_PPTX)
    print(f"Wrote {OUT_PPTX} ({len(deck.slides)} slide)")


if __name__ == "__main__":
    main()
