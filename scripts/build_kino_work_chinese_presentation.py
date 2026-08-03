#!/usr/bin/env python3
"""Build the Chinese whole-work presentation for KiNO / Kino-Fail.

The deck is intentionally receiver-oriented:

* problem and intuition precede terminology;
* complex evidence is introduced in small visual steps;
* all figure imagery is raster content from Isaac Sim, benchmark sensor
  records, measured data charts, or the current paper figures;
* the palette follows KiNO-Paper/figures/color.jpg.

The generated slide artwork never uses SVG.
"""

from __future__ import annotations

import json
import math
import shutil
from pathlib import Path

import matplotlib
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.font_manager import FontProperties
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PAPER = Path("/home/eureka/KiNO-Paper")
FIG = PAPER / "figures"
OUT_DIR = ROOT / "paper"
ASSET_DIR = OUT_DIR / "kino_work_cn_assets"
OUT_PPTX = OUT_DIR / "KiNO_work_chinese_presentation.pptx"

SW = 13.333
SH = 7.5

FONT_CN = "Noto Sans CJK SC"
FONT_CN_SERIF = "Noto Serif CJK SC"
FONT_EN = "Nimbus Sans"
FONT_CN_REG_FILE = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
FONT_CN_BOLD_FILE = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc")

# Palette copied from KiNO-Paper/figures/color.jpg.
TEAL = "2C9CA0"
PINK = "D77A7E"
RED = "C81223"
BLUE = "2665EA"
GRAY = "999999"
LIGHT_PINK = "F0C5C9"
BLACK = "000000"
WHITE = "FFFFFF"

INK = "17212B"
MUTED = "68727D"
LINE = "DDE2E6"
LIGHT = "F6F8F9"
TEAL_LIGHT = "E5F4F4"
BLUE_LIGHT = "EAF0FE"
RED_LIGHT = "FBE9EB"
PINK_LIGHT = "FAEFF0"
GRAY_LIGHT = "F0F1F2"

REPORT = ROOT / "outputs/eval/unified_moe_v3_publication_audit/report.json"
FIG2_ROOT = ROOT / "outputs/kinofail_realistic/fig2_registered_multiview_v1"
SCALE_ROOT = ROOT / "outputs/kinofail_realistic/corpus_scale_v7"
DEMO = ROOT / "outputs/demos/indoor_adhesion_icra"

OP_NAMES = {
    "O1": ("低摩擦", "脚底更容易打滑"),
    "O2": ("软地面", "脚会下陷"),
    "O3": ("地面塌陷", "支撑突然消失"),
    "O4": ("脚底粘附", "脚被地面粘住"),
    "O5": ("负载变化", "重量和重心改变"),
    "O6": ("侧向推力", "突然受到外力"),
    "O7": ("外观与物理错配", "看起来一样，实际更滑"),
    "O8": ("透明障碍", "相机难看见但会碰撞"),
    "O9": ("底盘托底", "横梁卡住机腹"),
    "O10": ("驱动力衰减", "电机出力受限"),
    "O11": ("身体感知偏差", "IMU 偏置与延迟"),
}

METHOD_CN = {
    "fixed_vision": "只看画面",
    "fixed_proprio": "只听身体",
    "fixed_joint": "直接拼接",
    "late_average": "后期平均",
    "legacy_class_rule": "旧规则",
    "random_route": "随机选择",
    "learned_router": "KiNO",
    "oracle_route": "理想上限",
}

METHOD_COLOR = {
    "fixed_vision": BLUE,
    "fixed_proprio": RED,
    "fixed_joint": GRAY,
    "late_average": PINK,
    "legacy_class_rule": LIGHT_PINK,
    "random_route": "BFC4C8",
    "learned_router": TEAL,
    "oracle_route": BLACK,
}


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def tint(hex6: str, alpha: float = 0.86) -> str:
    base = tuple(int(hex6[i : i + 2], 16) for i in (0, 2, 4))
    mixed = tuple(round(alpha * 255 + (1 - alpha) * c) for c in base)
    return "".join(f"{x:02X}" for x in mixed)


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def verify_sources() -> dict:
    """Verify the presentation's headline evidence against the publication audit."""

    data = read_json(REPORT)
    assert data["status"] == "complete_retrospective_publication_audit"
    assert data["counts"] == {
        "c2_development_cases": 60,
        "c2_development_samples": 360,
        "c2_test_cases": 30,
        "c2_test_samples": 180,
        "c4_cases": 75,
        "scale_material_families": 18,
        "scale_pairs": 988,
        "scale_samples": 5928,
        "scale_scene_instances": 9,
    }
    ours = data["cross_battery_summary"]["learned_router"]
    late = data["cross_battery_summary"]["late_average"]
    assert math.isclose(ours["equal_battery_macro_mean"], 0.975304279563267)
    assert math.isclose(ours["worst_battery_mean"], 0.9694974480154229)
    assert math.isclose(late["equal_battery_macro_mean"], 0.9722571274001941)
    assert math.isclose(late["worst_battery_mean"], 0.9600698103559437)
    route = data["decision_critical_route_audit"]["cells"]
    assert route["T2_vision_decisive"]["decision_critical_prediction_rows"] == 243
    assert route["T3_proprio_decisive"]["decision_critical_prediction_rows"] == 226
    assert math.isclose(route["T2_vision_decisive"]["evidence_route_fidelity"], 0.9300411522633745)
    assert math.isclose(route["T3_proprio_decisive"]["evidence_route_fidelity"], 1.0)
    strict = data["strict_group_certificates"]
    assert strict["scale"]["strict_all_training_seeds_and_group_members_correct"] == 295
    assert strict["conflict"]["strict_all_training_seeds_and_group_members_correct"] == 27
    c4 = data["c4_action_audit"]["methods"]["learned_router"]
    assert c4["release_cases"] == 15
    assert c4["release_precision"] == 1.0
    assert math.isclose(c4["closed_loop_vs_always_fallback"]["mean_terminal_cost_delta"], -5.934020734826723)
    assert math.isclose(c4["closed_loop_vs_always_fallback"]["mean_success_delta"], 0.2)
    return data


def set_mpl_style() -> FontProperties:
    matplotlib.use("Agg")
    prop = FontProperties(fname=str(FONT_CN_REG_FILE))
    matplotlib.rcParams.update(
        {
            "font.family": prop.get_name(),
            "font.sans-serif": [prop.get_name(), "DejaVu Sans"],
            "axes.unicode_minus": False,
            "axes.edgecolor": f"#{INK}",
            "axes.labelcolor": f"#{INK}",
            "xtick.color": f"#{MUTED}",
            "ytick.color": f"#{MUTED}",
            "text.color": f"#{INK}",
            "figure.facecolor": "white",
            "axes.facecolor": "white",
        }
    )
    return prop


def save_figure(fig: plt.Figure, path: Path) -> Path:
    fig.savefig(path, dpi=240, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_cross_battery_chart(data: dict) -> Path:
    rows = data["cross_battery_summary"]
    fig, ax = plt.subplots(figsize=(8.4, 5.0))
    ax.set_xlim(0.12, 1.015)
    ax.set_ylim(0.68, 1.015)
    ax.axvspan(0.95, 1.015, color=f"#{TEAL_LIGHT}", zorder=0)
    ax.axhspan(0.95, 1.015, color=f"#{TEAL_LIGHT}", zorder=0)
    ax.grid(color=f"#{LINE}", linewidth=0.8, alpha=0.9)
    order = [
        "fixed_vision",
        "fixed_proprio",
        "fixed_joint",
        "late_average",
        "legacy_class_rule",
        "random_route",
        "learned_router",
        "oracle_route",
    ]
    offsets = {
        "fixed_vision": (8, 2),
        "fixed_proprio": (-8, 8),
        "fixed_joint": (-28, 10),
        "late_average": (-12, -20),
        "legacy_class_rule": (-40, -16),
        "random_route": (-28, 8),
        "learned_router": (-28, 10),
        "oracle_route": (-48, -20),
    }
    for key in order:
        row = rows[key]
        x = row["scale_scene_and_material_balanced_accuracy"]["mean"]
        y = row["c2_conflict_balanced_accuracy"]["mean"]
        size = 230 if key == "learned_router" else 115
        marker = "*" if key == "learned_router" else ("X" if key == "oracle_route" else "o")
        edge = BLACK if key in {"learned_router", "oracle_route"} else WHITE
        ax.scatter(
            [x],
            [y],
            s=size,
            c=[f"#{METHOD_COLOR[key]}"],
            marker=marker,
            edgecolors=f"#{edge}",
            linewidths=1.2,
            zorder=4,
        )
        ax.annotate(
            METHOD_CN[key],
            (x, y),
            xytext=offsets[key],
            textcoords="offset points",
            fontsize=12 if key == "learned_router" else 10.5,
            weight="bold" if key == "learned_router" else "normal",
            color=f"#{METHOD_COLOR[key] if key != 'oracle_route' else BLACK}",
        )
    ax.set_xlabel("11 类全量测试：答对率", fontsize=12.5, weight="bold")
    ax.set_ylabel("4 类冲突测试：答对率", fontsize=12.5, weight="bold")
    ax.set_xticks([0.2, 0.4, 0.6, 0.8, 1.0])
    ax.set_yticks([0.7, 0.8, 0.9, 1.0])
    ax.text(
        0.985,
        0.695,
        "越靠右上越稳",
        ha="right",
        va="bottom",
        color=f"#{TEAL}",
        fontsize=11,
        weight="bold",
    )
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.tight_layout()
    return save_figure(fig, ASSET_DIR / "cross_battery_cn.png")


def make_route_chart(data: dict) -> Path:
    cells = data["decision_critical_route_audit"]["cells"]
    vals = [
        cells["T2_vision_decisive"]["evidence_route_fidelity"],
        cells["T3_proprio_decisive"]["evidence_route_fidelity"],
    ]
    labels = ["T2：该信画面", "T3：该信身体"]
    colors = [BLUE, RED]
    fig, ax = plt.subplots(figsize=(8.2, 2.65))
    y = np.arange(2)
    ax.barh(y, [1, 1], color=f"#{GRAY_LIGHT}", height=0.52)
    ax.barh(y, vals, color=[f"#{c}" for c in colors], height=0.52)
    for i, v in enumerate(vals):
        ax.text(v - 0.02, i, f"{v*100:.1f}%", ha="right", va="center", color="white", fontsize=18, weight="bold")
    ax.set_yticks(y, labels=labels, fontsize=13)
    ax.set_xlim(0, 1.02)
    ax.set_xticks([])
    ax.invert_yaxis()
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.set_title("专家意见冲突时，路由是否选对了证据来源？", fontsize=15, weight="bold", pad=12)
    fig.tight_layout()
    return save_figure(fig, ASSET_DIR / "route_fidelity_cn.png")


def make_scene_chart(data: dict) -> Path:
    scene = data["nine_scene_loso"]
    ours = scene["learned_router"]["scene_balanced_accuracy"]
    late = scene["late_average"]["scene_balanced_accuracy"]
    keys = list(ours)
    labels = ["野外-河道", "野外-坡地", "野外-林径", "室内-卧室", "室内-厨房", "室内-客厅", "室内-办公1", "室内-办公2", "室内-办公3"]
    x = np.arange(len(keys))
    fig, ax = plt.subplots(figsize=(9.1, 4.15))
    ax.plot(x, [ours[k] for k in keys], "-o", color=f"#{TEAL}", lw=2.6, ms=7, label="KiNO")
    ax.plot(x, [late[k] for k in keys], "-s", color=f"#{PINK}", lw=2.1, ms=6, label="后期平均")
    ax.fill_between(x, [late[k] for k in keys], [ours[k] for k in keys], color=f"#{TEAL}", alpha=0.08)
    ax.set_ylim(0.72, 1.015)
    ax.set_xticks(x, labels=labels, rotation=25, ha="right", fontsize=9.2)
    ax.set_ylabel("留出该场景后的答对率", fontsize=11.5, weight="bold")
    ax.grid(axis="y", color=f"#{LINE}", lw=0.8)
    ax.legend(frameon=False, ncol=2, loc="lower left")
    ax.text(
        8.05,
        0.997,
        f"9 场景平均：KiNO {scene['learned_router']['scene_mean']:.3f}",
        ha="right",
        va="top",
        color=f"#{TEAL}",
        fontsize=10.5,
        weight="bold",
    )
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig.tight_layout()
    return save_figure(fig, ASSET_DIR / "scene_generalization_cn.png")


def make_cost_asymmetry_chart() -> Path:
    matrix = np.array([[0, 1], [4, 1]], dtype=float)
    cmap = LinearSegmentedColormap.from_list("cost", [f"#{WHITE}", f"#{LIGHT_PINK}", f"#{RED}"])
    fig, ax = plt.subplots(figsize=(5.25, 4.1))
    ax.imshow(matrix, cmap=cmap, vmin=0, vmax=6)
    for i in range(2):
        for j in range(2):
            val = int(matrix[i, j])
            ax.text(j, i, str(val), ha="center", va="center", fontsize=25, weight="bold", color="white" if val >= 4 else "black")
    ax.set_xticks([0, 1], labels=["抬高脚步", "后退脱离"], fontsize=12)
    ax.set_yticks([0, 1], labels=["软地面", "脚底粘附"], fontsize=12)
    ax.set_xlabel("执行的动作", fontsize=12, weight="bold")
    ax.set_ylabel("真实原因", fontsize=12, weight="bold")
    ax.set_title("同一个动作，代价会因原因而反转", fontsize=14.5, weight="bold", pad=10)
    ax.tick_params(length=0)
    for spine in ax.spines.values():
        spine.set_visible(False)
    fig.tight_layout()
    return save_figure(fig, ASSET_DIR / "cost_asymmetry_cn.png")


def make_recovery_chart(data: dict) -> Path:
    pair = data["c4_action_audit"]["learned_router_pairwise"]
    labels = ["相对原地等待", "相对后期平均", "相对直接拼接"]
    keys = ["fixed_vision", "late_average", "fixed_joint"]
    # "fixed_vision" has the same zero-release action vector as the registered hold.
    gains = [-pair[k]["mean_terminal_cost_delta"] for k in keys]
    fig, ax = plt.subplots(figsize=(7.2, 3.6))
    y = np.arange(3)
    ax.barh(y, gains, color=[f"#{TEAL}", f"#{TEAL}", f"#{TEAL}"], height=0.54)
    for i, v in enumerate(gains):
        ax.text(v + 0.08, i, f"代价下降 {v:.2f}", va="center", fontsize=12.5, weight="bold", color=f"#{TEAL}")
    ax.set_yticks(y, labels=labels, fontsize=11.5)
    ax.set_xlim(0, 6.6)
    ax.set_xticks([])
    ax.invert_yaxis()
    ax.set_title("75 组相同起点的直接恢复试验", fontsize=15, weight="bold", pad=10)
    for spine in ax.spines.values():
        spine.set_visible(False)
    fig.tight_layout()
    return save_figure(fig, ASSET_DIR / "direct_recovery_cn.png")


def make_sensor_chart() -> tuple[Path, list[Path]]:
    manifests = sorted((SCALE_ROOT / "pilot/wild/O8_invisible_collider").glob("*_anomaly/manifest.json"))
    if not manifests:
        raise FileNotFoundError("No validated O8 scale manifest")
    manifest_path = manifests[0]
    manifest = read_json(manifest_path)
    rgb_rows = manifest["artifacts"]["rgb_frames"]
    frame_indices = np.linspace(0, len(rgb_rows) - 1, 5, dtype=int)
    frames = [manifest_path.parent / rgb_rows[int(i)]["path"] for i in frame_indices]

    proprio_path = manifest_path.parent / manifest["artifacts"]["proprio"]["path"]
    packet = np.load(proprio_path)
    t = packet["timestamp_s"]
    x = packet["features"]
    names = [str(v) for v in packet["feature_names"]]
    picks = {
        "身体倾斜": names.index("tilt_rad"),
        "滑移比例": names.index("slip_ratio"),
        "机身高度": names.index("base_height_m"),
        "支撑比例": names.index("support_ratio"),
    }
    fig, axes = plt.subplots(2, 2, figsize=(8.7, 3.7), sharex=True)
    cols = [RED, BLUE, TEAL, PINK]
    for ax, (label, idx), col in zip(axes.ravel(), picks.items(), cols, strict=True):
        ax.plot(t, x[:, idx], color=f"#{col}", lw=1.8)
        ax.set_title(label, fontsize=10.5, weight="bold")
        ax.grid(color=f"#{LINE}", lw=0.65)
        ax.tick_params(labelsize=8)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
    axes[-1, 0].set_xlabel("时间（秒）", fontsize=9.5)
    axes[-1, 1].set_xlabel("时间（秒）", fontsize=9.5)
    fig.tight_layout()
    return save_figure(fig, ASSET_DIR / "body_packet_cn.png"), frames


def crop_paper_figure_halves() -> tuple[Path, Path]:
    src = Image.open(FIG / "fig1_bidirectional_conflict.png").convert("RGB")
    w, h = src.size
    split = w // 2
    left = src.crop((0, 0, split + 10, h))
    right = src.crop((split - 10, 0, w, h))
    p1 = ASSET_DIR / "fig1_t2_half.png"
    p2 = ASSET_DIR / "fig1_t3_half.png"
    left.save(p1, optimize=True)
    right.save(p2, optimize=True)
    return p1, p2


def material_randomization_paths() -> tuple[list[Path], list[Path]]:
    manifest_path = (
        SCALE_ROOT
        / "pilot/wild/O7_visual_remap/cf_f9f8f122b1bcd4645700_anomaly/manifest.json"
    )
    manifest = read_json(manifest_path)
    paths = []
    for view in ("primary", "swap_01", "swap_02"):
        rows = manifest["artifacts"]["rgb_views"][view]
        idx = min(len(rows) - 1, max(0, int(len(rows) * 0.65)))
        paths.append(manifest_path.parent / rows[idx]["path"])
    matched_visuals = [FIG2_ROOT / f"O{i}/01_go2_front.png" for i in (1, 2, 3)]
    return paths, matched_visuals


def prepare_assets(data: dict) -> dict[str, Path | list[Path]]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    set_mpl_style()
    t2, t3 = crop_paper_figure_halves()
    body_chart, sensor_frames = make_sensor_chart()
    material_paths, same_look_paths = material_randomization_paths()
    return {
        "cross_battery": make_cross_battery_chart(data),
        "route": make_route_chart(data),
        "scene": make_scene_chart(data),
        "cost": make_cost_asymmetry_chart(),
        "recovery": make_recovery_chart(data),
        "body_chart": body_chart,
        "sensor_frames": sensor_frames,
        "t2": t2,
        "t3": t3,
        "materials": material_paths,
        "same_look": same_look_paths,
    }


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 20,
    color: str = INK,
    bold: bool = False,
    font: str = FONT_CN,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.03,
    line_spacing: float | None = None,
    italic: bool = False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.text = text
    for paragraph in tf.paragraphs:
        paragraph.alignment = align
        paragraph.space_after = Pt(0)
        paragraph.space_before = Pt(0)
        if line_spacing is not None:
            paragraph.line_spacing = line_spacing
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = rgb(color)
    return box


def add_runs(
    slide,
    runs: list[tuple[str, dict]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 20,
    color: str = INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for content, opts in runs:
        r = p.add_run()
        r.text = content
        r.font.name = opts.get("font", FONT_CN)
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", False)
        r.font.italic = opts.get("italic", False)
        r.font.color.rgb = rgb(opts.get("color", color))
    return box


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str = LINE,
    width: float = 1.0,
    rounded: bool = True,
    transparency: int = 0,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *, color: str = LINE, width: float = 1.0):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_badge(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    *,
    fill: str = TEAL_LIGHT,
    color: str = TEAL,
    size: float = 10.5,
):
    add_box(slide, x, y, w, 0.34, fill=fill, line=fill, rounded=True)
    add_text(
        slide,
        text,
        x + 0.06,
        y + 0.005,
        w - 0.12,
        0.32,
        size=size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_picture_crop(
    slide,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    focus_x: float = 0.5,
    focus_y: float = 0.5,
    border: str | None = None,
    border_width: float = 1.0,
):
    with Image.open(path) as image:
        iw, ih = image.size
    src_ar = iw / ih
    dst_ar = w / h
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if src_ar > dst_ar:
        visible = dst_ar / src_ar
        excess = 1.0 - visible
        left = max(0.0, min(excess, excess * focus_x))
        pic.crop_left = left
        pic.crop_right = excess - left
    elif src_ar < dst_ar:
        visible = src_ar / dst_ar
        excess = 1.0 - visible
        top = max(0.0, min(excess, excess * focus_y))
        pic.crop_top = top
        pic.crop_bottom = excess - top
    if border:
        pic.line.color.rgb = rgb(border)
        pic.line.width = Pt(border_width)
    return pic


def add_picture_contain(
    slide,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    align_x: float = 0.5,
    align_y: float = 0.5,
):
    """Fit a raster image inside a box without cropping or distorting it."""

    with Image.open(path) as image:
        iw, ih = image.size
    src_ar = iw / ih
    dst_ar = w / h
    if src_ar >= dst_ar:
        draw_w = w
        draw_h = w / src_ar
        draw_x = x
        draw_y = y + (h - draw_h) * align_y
    else:
        draw_h = h
        draw_w = h * src_ar
        draw_x = x + (w - draw_w) * align_x
        draw_y = y
    return slide.shapes.add_picture(
        str(path),
        Inches(draw_x),
        Inches(draw_y),
        Inches(draw_w),
        Inches(draw_h),
    )


def add_image_card(
    slide,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    label: str | None = None,
    label_fill: str = BLACK,
    label_color: str = WHITE,
    focus_x: float = 0.5,
    focus_y: float = 0.5,
    border: str = WHITE,
):
    add_box(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, fill=WHITE, line=border, rounded=False)
    add_picture_crop(slide, path, x, y, w, h, focus_x=focus_x, focus_y=focus_y)
    if label:
        label_w = min(w - 0.22, max(1.08, len(label) * 0.23 + 0.34))
        add_box(slide, x + 0.10, y + 0.10, label_w, 0.4, fill=label_fill, line=label_fill, rounded=True, transparency=7)
        add_text(
            slide,
            label,
            x + 0.16,
            y + 0.115,
            label_w - 0.12,
            0.36,
            size=12,
            color=label_color,
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
        )


def add_metric(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    value: str,
    label: str,
    *,
    accent: str = TEAL,
    fill: str = WHITE,
    detail: str = "",
):
    add_box(slide, x, y, w, h, fill=fill, line=tint(accent, 0.75), width=1.1, rounded=True)
    add_box(slide, x, y, 0.08, h, fill=accent, line=accent, rounded=False)
    add_text(slide, value, x + 0.25, y + 0.12, w - 0.4, 0.48, size=25, color=accent, bold=True)
    add_text(slide, label, x + 0.25, y + 0.63, w - 0.4, 0.34, size=12.5, color=INK, bold=True)
    if detail:
        add_text(slide, detail, x + 0.25, y + 0.98, w - 0.4, h - 1.03, size=9.5, color=MUTED)


def add_bullets(
    slide,
    items: list[tuple[str, str]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 17,
    color: str = INK,
    bullet_color: str = TEAL,
    gap: float = 0.14,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for idx, (head, body) in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(gap * 72)
        p.line_spacing = 1.06
        r = p.add_run()
        r.text = "●  "
        r.font.name = FONT_CN
        r.font.size = Pt(size - 2)
        r.font.color.rgb = rgb(bullet_color)
        r = p.add_run()
        r.text = head
        r.font.name = FONT_CN
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = rgb(color)
        if body:
            r = p.add_run()
            r.text = f"：{body}"
            r.font.name = FONT_CN
            r.font.size = Pt(size)
            r.font.color.rgb = rgb(color)
    return box


def add_slide_header(slide, section: str, title: str, number: int, *, source: str = ""):
    add_text(slide, section, 0.67, 0.28, 3.0, 0.25, size=10.5, color=TEAL, bold=True)
    add_text(slide, title, 0.67, 0.60, 11.95, 0.56, size=27.5, color=INK, bold=True)
    add_line(slide, 0.67, 7.06, 12.66, 7.06, color=LINE, width=0.7)
    if source:
        add_text(slide, source, 0.68, 7.13, 10.8, 0.18, size=7.8, color=GRAY)
    add_text(slide, f"{number:02d}", 12.10, 7.12, 0.55, 0.18, size=8.5, color=GRAY, align=PP_ALIGN.RIGHT)


def add_takeaway(slide, text: str, *, y: float = 6.42, fill: str = TEAL_LIGHT, color: str = TEAL):
    add_box(slide, 0.67, y, 11.99, 0.45, fill=fill, line=fill, rounded=True)
    add_text(slide, text, 0.90, y + 0.04, 11.55, 0.36, size=15.3, color=color, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def new_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    return slide


def operator_path(operator: str, view: str) -> Path:
    mapping = {
        "front": "01_go2_front.png",
        "contact": "02_contact_left.png",
        "contact_right": "03_contact_right.png",
        "overview": "04_overview.png",
    }
    return FIG2_ROOT / operator / mapping[view]


def add_operator_card(slide, op: str, x: float, y: float, w: float, h: float):
    title, desc = OP_NAMES[op]
    add_box(slide, x, y, w, h, fill=WHITE, line=LINE, width=1.0, rounded=True)
    image_h = h * 0.63
    add_picture_crop(slide, operator_path(op, "overview"), x + 0.08, y + 0.08, w - 0.16, image_h, focus_y=0.52)
    inset_w = min(1.30, w * 0.40)
    inset_h = inset_w * 0.57
    add_picture_crop(
        slide,
        operator_path(op, "contact"),
        x + w - inset_w - 0.14,
        y + 0.08 + image_h - inset_h - 0.08,
        inset_w,
        inset_h,
        border=TEAL,
        border_width=2.0,
    )
    add_badge(slide, op, x + 0.16, y + 0.16, 0.52, fill=BLACK, color=WHITE, size=11)
    add_text(slide, title, x + 0.14, y + image_h + 0.19, w - 0.28, 0.34, size=15.5, color=INK, bold=True)
    add_text(slide, desc, x + 0.14, y + image_h + 0.56, w - 0.28, h - image_h - 0.62, size=10.8, color=MUTED)


def add_four_answer_card(slide, index: str, title: str, body: str, image: Path, x: float, accent: str):
    add_box(slide, x, 1.62, 2.84, 4.62, fill=WHITE, line=LINE, width=1.0, rounded=True)
    add_picture_crop(slide, image, x + 0.09, 1.71, 2.66, 1.76, focus_y=0.55)
    add_badge(slide, index, x + 0.18, 1.81, 0.48, fill=accent, color=WHITE, size=10.5)
    add_text(slide, title, x + 0.20, 3.72, 2.44, 0.42, size=18.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.23, 4.30, 2.38, 1.25, size=13, color=INK, align=PP_ALIGN.CENTER, line_spacing=1.08)


def build_deck(data: dict, assets: dict[str, Path | list[Path]]) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    # 01 — cover
    slide = new_slide(prs)
    add_badge(slide, "KiNO · 中文工作汇报", 0.72, 0.52, 2.25, fill=TEAL_LIGHT, color=TEAL)
    add_text(slide, "Feel It, See It,\nRecover", 0.72, 1.18, 5.72, 1.48, size=37, color=INK, bold=True, font=FONT_EN)
    add_text(slide, "让四足机器人看懂“为什么会失败”\n再决定如何恢复", 0.75, 2.95, 5.45, 1.15, size=25, color=INK, bold=True, line_spacing=1.05)
    add_text(slide, "跨视觉与身体感知的异常归因方法 + Kino-Fail 真实仿真基准", 0.76, 4.40, 5.36, 0.72, size=16, color=MUTED)
    add_runs(
        slide,
        [
            ("同样是“走不动”  ", {"color": MUTED, "size": 15}),
            ("原因不同  ", {"color": RED, "size": 15, "bold": True}),
            ("动作也必须不同", {"color": TEAL, "size": 15, "bold": True}),
        ],
        0.76,
        5.55,
        5.48,
        0.45,
    )
    add_text(slide, "研究组汇报 · 2026.07", 0.76, 6.70, 3.5, 0.24, size=10.5, color=GRAY)
    add_picture_crop(slide, DEMO / "multiview_peak/08_room_wide.png", 6.65, 0.50, 6.02, 3.20, focus_y=0.52)
    add_picture_crop(slide, operator_path("O8", "overview"), 6.65, 3.84, 2.90, 2.65, focus_y=0.50)
    add_picture_crop(slide, operator_path("O9", "overview"), 9.72, 3.84, 2.95, 2.65, focus_y=0.50)
    add_badge(slide, "脚底粘附", 6.86, 0.72, 1.20, fill=BLACK, color=WHITE)
    add_badge(slide, "透明障碍", 6.86, 4.06, 1.20, fill=BLACK, color=WHITE)
    add_badge(slide, "底盘托底", 9.93, 4.06, 1.20, fill=BLACK, color=WHITE)
    add_text(slide, "01", 12.10, 7.12, 0.55, 0.18, size=8.5, color=GRAY, align=PP_ALIGN.RIGHT)

    # 02 — problem
    slide = new_slide(prs)
    add_slide_header(slide, "01 · 问题", "机器人卡住了——真正重要的是：为什么？", 2, source="Isaac Sim registered multiview cases")
    cards = [
        ("O2", "软地面", "需要抬高脚步"),
        ("O4", "脚被粘住", "需要后退脱离"),
        ("O8", "碰到透明障碍", "需要停住重规划"),
        ("O9", "底盘被横梁托住", "需要抬身并减速"),
    ]
    for i, (op, title, action) in enumerate(cards):
        x = 0.70 + i * 3.05
        add_image_card(slide, operator_path(op, "overview"), x, 1.48, 2.82, 2.45, label=title)
        add_text(slide, action, x + 0.12, 4.15, 2.58, 0.40, size=15.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, "表面症状都可能是：速度下降、姿态异常或停滞", x + 0.20, 4.72, 2.42, 0.90, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "检测到“出问题”只是开始；恢复动作取决于对物理原因的判断。", y=6.17)

    # 03 — asymmetric consequence
    slide = new_slide(prs)
    add_slide_header(slide, "01 · 问题", "错误归因不只是答错标签，而是会执行错误动作", 3, source="Paper Fig. 5b; A4 forced-label intervention")
    add_image_card(slide, operator_path("O2", "contact"), 0.72, 1.48, 2.75, 2.10, label="软地面", label_fill=BLUE)
    add_image_card(slide, operator_path("O4", "contact"), 3.68, 1.48, 2.75, 2.10, label="脚底粘附", label_fill=RED)
    add_text(slide, "正确动作：抬高脚步", 0.92, 3.84, 2.34, 0.36, size=15, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "正确动作：后退脱离", 3.88, 3.84, 2.34, 0.36, size=15, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_picture_crop(slide, assets["cost"], 6.78, 1.30, 5.72, 4.35)
    add_metric(slide, 0.86, 4.54, 2.48, 1.27, "4", "粘附时误用抬脚的代价", accent=RED, fill=RED_LIGHT)
    add_metric(slide, 3.83, 4.54, 2.48, 1.27, "1", "软地面时保守后退的代价", accent=PINK, fill=PINK_LIGHT)
    add_takeaway(slide, "在这组已登记的代价模型中，危险方向比保守方向重 4 倍。", y=6.20, fill=RED_LIGHT, color=RED)

    # 04 — overall conflict
    slide = new_slide(prs)
    add_slide_header(slide, "02 · 核心发现", "眼睛和身体谁更可信？答案会随异常类型反转", 4, source="Current paper Fig. 1")
    add_picture_contain(slide, FIG / "fig1_bidirectional_conflict.png", 0.67, 1.35, 12.0, 4.66)
    add_box(slide, 0.78, 5.94, 5.72, 0.42, fill=BLUE_LIGHT, line=BLUE_LIGHT)
    add_text(slide, "T2：身体感觉完全相同 → 看场景", 0.96, 5.98, 5.35, 0.32, size=14.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 6.84, 5.94, 5.72, 0.42, fill=RED_LIGHT, line=RED_LIGHT)
    add_text(slide, "T3：画面完全相同 → 听身体", 7.02, 5.98, 5.35, 0.32, size=14.5, color=RED, bold=True, align=PP_ALIGN.CENTER)

    # 05 — T2 gradual explanation
    slide = new_slide(prs)
    add_slide_header(slide, "02 · 核心发现", "T2：在接触之前，场景外观提供决定性线索", 5, source="Exact byte-identical pre-contact body packet")
    add_picture_contain(slide, assets["t2"], 0.68, 1.35, 7.28, 4.82)
    add_badge(slide, "接触前判断", 8.35, 1.52, 1.45, fill=BLUE_LIGHT, color=BLUE)
    add_text(slide, "软地面与粘附地面\n看起来不同", 8.34, 2.12, 3.85, 0.85, size=23, color=BLUE, bold=True)
    add_bullets(
        slide,
        [
            ("身体输入完全相同", "21×19 个数逐字节一致"),
            ("视觉能区分原因", "一个需要抬脚，一个需要后退"),
            ("因此只听身体一定会猜", "不是模型容量问题，而是输入没有信息"),
        ],
        8.34,
        3.25,
        4.10,
        2.48,
        size=15.3,
        bullet_color=BLUE,
    )
    add_takeaway(slide, "这组实验不是“视觉更强”，而是证明这里必须使用视觉。", y=6.35, fill=BLUE_LIGHT, color=BLUE)

    # 06 — T3 gradual explanation
    slide = new_slide(prs)
    add_slide_header(slide, "02 · 核心发现", "T3：接触之后，身体反应揭示画面看不到的物理差异", 6, source="Pixel-identical five-frame sequence through the encounter boundary")
    add_picture_contain(slide, assets["t3"], 5.37, 1.35, 7.28, 4.82)
    add_badge(slide, "接触后判断", 0.80, 1.52, 1.45, fill=RED_LIGHT, color=RED)
    add_text(slide, "低摩擦与透明障碍\n可以给出完全相同的画面", 0.79, 2.12, 4.18, 0.92, size=23, color=RED, bold=True)
    add_bullets(
        slide,
        [
            ("五帧画面像素一致", "相机无法判断哪一种物理原因"),
            ("身体反应明显不同", "一个持续滑动，一个产生碰撞冲击"),
            ("因此只看画面一定会漏", "需要接触后的加速度和速度历史"),
        ],
        0.80,
        3.30,
        4.10,
        2.35,
        size=15.3,
        bullet_color=RED,
    )
    add_takeaway(slide, "这组实验反向证明：有些异常必须依靠身体感知。", y=6.35, fill=RED_LIGHT, color=RED)

    # 07 — four questions / contributions
    slide = new_slide(prs)
    add_slide_header(slide, "03 · 整体工作", "整个工作围绕四个简单问题展开", 7)
    answer_cards = [
        ("1", "问题是什么？", "把导航失败从“是否异常”推进到“是什么物理原因”。", operator_path("O8", "overview"), BLUE),
        ("2", "如何可控地测？", "用 11 个仿真算子，构造可重复的反事实异常。", operator_path("O3", "overview"), PINK),
        ("3", "如何判断？", "让视觉、身体和联合专家竞争，再学习该信谁。", operator_path("O4", "overview"), TEAL),
        ("4", "判断有用吗？", "把原因连接到恢复动作，并测量真实物理后果。", DEMO / "nominal/review_t07.5s.png", RED),
    ]
    for i, (num, title, body, image, accent) in enumerate(answer_cards):
        add_four_answer_card(
            slide,
            num,
            title,
            body,
            image,
            0.70 + i * 3.04,
            accent,
        )
    add_takeaway(slide, "问题定义、基准、方法和动作后果组成一条完整证据链。", y=6.43)

    # 08 — dataset overview
    slide = new_slide(prs)
    add_slide_header(slide, "04 · Kino-Fail", "一个专门测“异常原因能否被看懂”的真实仿真基准", 8, source="Registered design and frozen scale snapshot")
    add_picture_crop(slide, DEMO / "multiview_peak/01_room_overview_left.png", 0.70, 1.40, 3.80, 2.55)
    add_picture_crop(slide, operator_path("O6", "overview"), 4.75, 1.40, 3.80, 2.55)
    add_picture_crop(slide, operator_path("O1", "overview"), 8.80, 1.40, 3.80, 2.55)
    add_badge(slide, "生活 / 室内", 0.90, 1.60, 1.30, fill=BLACK, color=WHITE)
    add_badge(slide, "生产 / 设施", 4.95, 1.60, 1.30, fill=BLACK, color=WHITE)
    add_badge(slide, "野外", 9.00, 1.60, 0.82, fill=BLACK, color=WHITE)
    add_metric(slide, 0.76, 4.35, 2.28, 1.42, "11", "仿真异常算子", accent=TEAL, detail="覆盖地面、接触、外力、机器人和传感器")
    add_metric(slide, 3.26, 4.35, 2.28, 1.42, "988", "冻结反事实对", accent=BLUE, detail="每对从干净重置开始，只改一个物理因素")
    add_metric(slide, 5.76, 4.35, 2.28, 1.42, "1,976", "物理仿真 episode", accent=PINK, detail="每对包含正常与异常两侧")
    add_metric(slide, 8.26, 4.35, 2.28, 1.42, "5,928", "同步视觉记录", accent=RED, detail="每个 episode 含 3 种独立外观")
    add_metric(slide, 10.76, 4.35, 1.84, 1.42, "18", "材质家族", accent=GRAY, detail="外观与物理分开变化")
    add_takeaway(slide, "RGB 来自 Go2 机身前视 RTX 相机；解释性低视角只用于展示，不喂给模型。", y=6.22)

    # 09 — operators 1–6
    slide = new_slide(prs)
    add_slide_header(slide, "04 · Kino-Fail", "11 种异常算子（1/2）：地面、接触与外力", 9, source="Registered same-state Go2-front + low-contact views")
    layout = [(0.70, 1.35), (4.30, 1.35), (7.90, 1.35), (0.70, 4.05), (4.30, 4.05), (7.90, 4.05)]
    for op, (x, y) in zip(["O1", "O2", "O3", "O4", "O5", "O6"], layout, strict=True):
        add_operator_card(slide, op, x, y, 3.38, 2.42)
    add_box(slide, 11.50, 1.35, 1.10, 5.12, fill=TEAL_LIGHT, line=TEAL_LIGHT)
    add_text(slide, "每个卡片\n\n大图：\n机器人所处场景\n\n小图：\n脚边/接触证据", 11.66, 1.80, 0.78, 4.15, size=12.2, color=TEAL, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.06)

    # 10 — operators 7–11
    slide = new_slide(prs)
    add_slide_header(slide, "04 · Kino-Fail", "11 种异常算子（2/2）：错配、几何、机器人与传感器", 10, source="Registered same-state Go2-front + low-contact views")
    layout = [(0.75, 1.38), (4.35, 1.38), (7.95, 1.38), (2.55, 4.07), (6.15, 4.07)]
    for op, (x, y) in zip(["O7", "O8", "O9", "O10", "O11"], layout, strict=True):
        add_operator_card(slide, op, x, y, 3.38, 2.42)
    add_box(slide, 9.84, 4.07, 2.72, 2.42, fill=LIGHT, line=LINE)
    add_text(slide, "为什么需要这些算子？", 10.10, 4.38, 2.18, 0.42, size=17, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "真实失败并不都能从一张图直接看见。\n\n我们既测“外观会骗人”，也测“身体传感会骗人”。", 10.10, 4.98, 2.18, 1.13, size=13.1, color=INK, align=PP_ALIGN.CENTER, line_spacing=1.05)

    # 11 — appearance / physics independence
    slide = new_slide(prs)
    add_slide_header(slide, "04 · Kino-Fail", "关键设计：外观与物理分开随机化，避免模型记住贴图", 11, source="Frozen RGB view variants; same-state registered operator views")
    material_paths = assets["materials"]
    same_look_paths = assets["same_look"]
    add_text(slide, "同一种物理，可以长得不一样", 0.76, 1.35, 5.45, 0.38, size=18, color=TEAL, bold=True)
    for i, (p, lab) in enumerate(zip(material_paths, ["材质 A", "材质 B", "材质 C"], strict=True)):
        add_image_card(slide, p, 0.75 + i * 2.03, 1.82, 1.86, 1.50, label=lab)
    add_text(slide, "同一种外观，也可以承载不同物理", 6.93, 1.35, 5.45, 0.38, size=18, color=RED, bold=True)
    for i, (p, lab) in enumerate(zip(same_look_paths, ["低摩擦", "软地面", "会塌陷"], strict=True)):
        add_image_card(slide, p, 6.92 + i * 1.88, 1.82, 1.72, 1.50, label=lab)
    add_box(slide, 0.75, 3.72, 5.85, 2.10, fill=TEAL_LIGHT, line=TEAL_LIGHT)
    add_text(slide, "模型不能靠“看见某种纹理”就猜物理标签", 1.08, 4.06, 5.18, 0.46, size=19, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "每个物理 episode 额外渲染多个外观；训练时这些外观的总权重仍为 1，不虚增样本量。", 1.10, 4.74, 5.12, 0.74, size=13.2, color=INK, align=PP_ALIGN.CENTER)
    add_box(slide, 6.79, 3.72, 5.81, 2.10, fill=RED_LIGHT, line=RED_LIGHT)
    add_text(slide, "物理标签来自仿真算子，不来自贴图名字", 7.10, 4.06, 5.18, 0.46, size=19, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "摩擦、柔顺、碰撞、外力等由 Isaac Sim 物理层独立控制；材质只改变画面。", 7.14, 4.74, 5.08, 0.74, size=13.2, color=INK, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "目标是让模型学“证据关系”，而不是把某块地毯记成某种异常。", y=6.30)

    # 12 — observation packet
    slide = new_slide(prs)
    add_slide_header(slide, "05 · 输入", "每次判断只看两样东西：眼前几帧 + 身体一段反应", 12, source="Frozen O8 benchmark episode; Go2-front RGB and 19-channel packet")
    sensor_frames = assets["sensor_frames"]
    for i, p in enumerate(sensor_frames):
        add_picture_crop(slide, p, 0.72 + i * 1.42, 1.50, 1.28, 1.45, focus_y=0.58)
        add_text(slide, f"t{i+1}", 1.11 + i * 1.42, 2.72, 0.50, 0.20, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "眼睛：5 帧机器人前视画面", 0.78, 3.07, 6.75, 0.36, size=18, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, assets["body_chart"], 7.54, 1.38, 5.04, 3.52)
    add_text(slide, "身体：19 路原始信号 → 8 类易解释量", 7.73, 5.00, 4.64, 0.38, size=18, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 0.75, 4.05, 6.62, 1.80, fill=BLUE_LIGHT, line=BLUE_LIGHT)
    add_text(slide, "画面回答“前方是什么”", 1.08, 4.38, 5.96, 0.42, size=20, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "身体回答“接触以后发生了什么”\n包括倾斜、滑移、支撑、速度、加速度和用力程度", 1.06, 4.96, 6.00, 0.72, size=13.4, color=INK, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "不读取场景编号、算子真值、材质编号或仿真器内部标签。", y=6.27, fill=GRAY_LIGHT, color=MUTED)

    # 13 — counterfactual protocol
    slide = new_slide(prs)
    add_slide_header(slide, "06 · 实验原则", "反事实配对：同一个起点，只改变一个物理原因", 13, source="Indoor adhesion demo and registered reset-isolated protocol")
    add_image_card(slide, DEMO / "nominal/review_t04.8s.png", 0.75, 1.47, 4.05, 2.90, label="正常侧", label_fill=TEAL)
    add_image_card(slide, DEMO / "adhesion_peel/review_t04.8s.png", 8.53, 1.47, 4.05, 2.90, label="异常侧：脚底粘附", label_fill=RED)
    add_box(slide, 5.17, 1.90, 3.00, 2.10, fill=LIGHT, line=LINE, width=1.2)
    add_text(slide, "同一场景\n同一初始状态\n同一路线\n同一相机", 5.50, 2.16, 2.34, 1.30, size=17, color=INK, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.05)
    add_text(slide, "只切换一个算子", 5.33, 3.61, 2.67, 0.30, size=14.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 4.80, 2.92, 5.17, 2.92, color=TEAL, width=3)
    add_line(slide, 8.17, 2.92, 8.53, 2.92, color=RED, width=3)
    add_metric(slide, 0.98, 4.78, 3.38, 1.25, "每对都重置", "排除前一条轨迹的残留影响", accent=TEAL, fill=TEAL_LIGHT)
    add_metric(slide, 4.98, 4.78, 3.38, 1.25, "保存哈希", "场景、相机、边界和文件均可核验", accent=BLUE, fill=BLUE_LIGHT)
    add_metric(slide, 8.98, 4.78, 3.38, 1.25, "配对统计", "把“哪种方法更好”落在同一案例上", accent=RED, fill=RED_LIGHT)
    add_takeaway(slide, "我们测的是归因方法本身，而不是某次导航随机跑得更幸运。", y=6.27)

    # 14 — method
    slide = new_slide(prs)
    add_slide_header(slide, "07 · 方法", "KiNO 的直觉：三位专家 + 一位会看分歧的裁判", 14, source="Current paper Fig. 3 architecture")
    add_picture_contain(slide, FIG / "fig3_structured_router_legacy_boxes_backup.png", 0.73, 1.35, 6.55, 4.88)
    add_badge(slide, "视觉专家", 7.67, 1.52, 1.22, fill=BLUE_LIGHT, color=BLUE)
    add_text(slide, "只从画面判断", 9.05, 1.56, 2.85, 0.32, size=16, color=INK, bold=True)
    add_badge(slide, "身体专家", 7.67, 2.42, 1.22, fill=RED_LIGHT, color=RED)
    add_text(slide, "只从身体反应判断", 9.05, 2.46, 2.85, 0.32, size=16, color=INK, bold=True)
    add_badge(slide, "联合专家", 7.67, 3.32, 1.22, fill=PINK_LIGHT, color=PINK)
    add_text(slide, "直接同时使用两路输入", 9.05, 3.36, 3.12, 0.32, size=16, color=INK, bold=True)
    add_box(slide, 7.64, 4.17, 4.74, 1.43, fill=TEAL_LIGHT, line=TEAL, width=1.2)
    add_text(slide, "学习型路由器", 7.95, 4.44, 4.12, 0.36, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "观察三位专家的把握、分歧和不确定性，再选择这次该信谁。", 7.96, 4.94, 4.10, 0.52, size=13.7, color=INK, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "重点不是把两种输入拼在一起，而是显式学习“这次哪一路证据更有诊断力”。", y=6.30)

    # 15 — routing examples
    slide = new_slide(prs)
    add_slide_header(slide, "07 · 方法", "同一套模型，在两种冲突里走不同证据路线", 15)
    add_box(slide, 0.72, 1.42, 5.88, 4.86, fill=BLUE_LIGHT, line=BLUE_LIGHT)
    add_text(slide, "T2 · 接触前", 0.98, 1.68, 1.70, 0.38, size=19, color=BLUE, bold=True)
    add_image_card(slide, operator_path("O2", "front"), 0.98, 2.20, 2.38, 1.72, label="软地面")
    add_image_card(slide, operator_path("O4", "front"), 3.74, 2.20, 2.38, 1.72, label="脚底粘附")
    add_text(slide, "身体输入相同", 1.12, 4.18, 1.92, 0.34, size=14, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→ 选择视觉证据", 3.87, 4.18, 2.07, 0.34, size=16, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "画面中的软土/粘附膜语义\n决定下一步动作", 1.31, 4.86, 4.71, 0.75, size=16, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 6.75, 1.42, 5.86, 4.86, fill=RED_LIGHT, line=RED_LIGHT)
    add_text(slide, "T3 · 接触后", 7.01, 1.68, 1.70, 0.38, size=19, color=RED, bold=True)
    add_image_card(slide, operator_path("O7", "front"), 7.01, 2.20, 2.38, 1.72, label="低摩擦")
    add_image_card(slide, operator_path("O8", "front"), 9.77, 2.20, 2.38, 1.72, label="透明障碍")
    add_text(slide, "画面输入相同", 7.15, 4.18, 1.92, 0.34, size=14, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "→ 选择身体证据", 9.90, 4.18, 2.07, 0.34, size=16, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "滑移与碰撞产生的身体反应\n决定真实物理原因", 7.34, 4.86, 4.71, 0.75, size=16, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "路由结果随事件改变，而不是永远偏向某一种传感器。", y=6.42)

    # 16 — main result
    slide = new_slide(prs)
    add_slide_header(slide, "08 · 主要结果", "一套模型同时面对“全量异常”与“专门冲突”两种测试", 16, source="Publication audit: five checkpoints, no test refit")
    add_picture_contain(slide, assets["cross_battery"], 0.72, 1.33, 7.62, 4.87)
    add_box(slide, 8.64, 1.47, 3.83, 1.16, fill=TEAL_LIGHT, line=TEAL)
    add_text(slide, "97.5%", 8.96, 1.68, 1.45, 0.45, size=27, color=TEAL, bold=True)
    add_text(slide, "两套测试等权平均", 10.30, 1.72, 1.88, 0.36, size=13.2, color=INK, bold=True)
    add_box(slide, 8.64, 2.84, 3.83, 1.16, fill=BLUE_LIGHT, line=BLUE)
    add_text(slide, "96.9%", 8.96, 3.05, 1.45, 0.45, size=27, color=BLUE, bold=True)
    add_text(slide, "较差一套测试的成绩", 10.30, 3.09, 1.88, 0.36, size=13.2, color=INK, bold=True)
    add_box(slide, 8.64, 4.21, 3.83, 1.45, fill=LIGHT, line=LINE)
    add_text(slide, "不是每个单项冠军", 8.94, 4.45, 3.22, 0.36, size=17, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "但在两类证据规则反转时，\nKiNO 的整体最差表现最高。", 8.92, 4.91, 3.25, 0.58, size=13.1, color=INK, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "相对后期平均：等权平均 +0.3 个百分点，最差测试 +0.9 个百分点。", y=6.28)

    # 17 — route audit / selective risk
    slide = new_slide(prs)
    add_slide_header(slide, "08 · 主要结果", "方法不仅答对，还能解释“这次为什么信这一路”", 17, source="Decision-critical route audit; same-coverage risk")
    add_picture_contain(slide, assets["route"], 0.70, 1.30, 7.20, 2.46)
    add_image_card(slide, operator_path("O2", "front"), 0.86, 4.10, 1.70, 1.32, label="T2")
    add_image_card(slide, operator_path("O8", "front"), 2.82, 4.10, 1.70, 1.32, label="T3")
    add_text(slide, "243 行专家分歧", 4.75, 4.17, 2.58, 0.34, size=15, color=BLUE, bold=True)
    add_text(slide, "226 行专家分歧", 4.75, 4.68, 2.58, 0.34, size=15, color=RED, bold=True)
    add_box(slide, 8.22, 1.46, 4.22, 4.30, fill=LIGHT, line=LINE)
    add_text(slide, "同等覆盖率下的错误风险", 8.55, 1.78, 3.56, 0.42, size=19, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_metric(slide, 8.62, 2.55, 3.42, 1.18, "1.2%", "只处理最有把握的 75% 组", accent=TEAL, fill=TEAL_LIGHT)
    add_metric(slide, 8.62, 4.03, 3.42, 1.18, "7.2%", "全部处理时的两套测试平均风险", accent=BLUE, fill=BLUE_LIGHT)
    add_takeaway(slide, "路由是可审计的中间决定：我们能检查它在冲突时是否信对了证据。", y=6.28)

    # 18 — generalization
    slide = new_slide(prs)
    add_slide_header(slide, "08 · 主要结果", "换场景、换材质后，结论仍保持稳定", 18, source="Nine-scene LOSO; six held-out material families; strict group certificates")
    add_picture_contain(slide, assets["scene"], 0.68, 1.32, 7.72, 4.35)
    add_image_card(slide, DEMO / "multiview_peak/06_front_three_quarter.png", 8.67, 1.46, 1.77, 1.34, label="室内")
    add_image_card(slide, operator_path("O6", "overview"), 10.66, 1.46, 1.77, 1.34, label="设施")
    add_image_card(slide, operator_path("O1", "overview"), 8.67, 3.04, 3.76, 1.75, label="野外")
    add_metric(slide, 0.82, 5.70, 2.82, 0.84, "93.8%", "留一场景平均", accent=TEAL)
    add_metric(slide, 3.84, 5.70, 2.82, 0.84, "95.2%", "最弱测试材质", accent=BLUE)
    add_metric(slide, 6.86, 5.70, 2.82, 0.84, "295 / 329", "全种子整对全对", accent=PINK)
    add_metric(slide, 9.88, 5.70, 2.55, 0.84, "27 / 30", "冲突案例全对", accent=RED)

    # 19 — action consequence
    slide = new_slide(prs)
    add_slide_header(slide, "09 · 从归因到动作", "归因必须改变动作，也必须改变真实物理后果", 19, source="630 forced-label Isaac Go2 interventions")
    add_picture_contain(slide, FIG / "fig5_intervention_consequence.png", 0.69, 1.32, 7.50, 4.78)
    add_box(slide, 8.52, 1.50, 3.83, 1.13, fill=TEAL_LIGHT, line=TEAL)
    add_text(slide, "630", 8.84, 1.70, 1.10, 0.44, size=27, color=TEAL, bold=True)
    add_text(slide, "受控物理试验", 9.88, 1.74, 2.02, 0.34, size=15, color=INK, bold=True)
    add_box(slide, 8.52, 2.83, 3.83, 1.13, fill=BLUE_LIGHT, line=BLUE)
    add_text(slide, "9 / 9", 8.84, 3.03, 1.10, 0.44, size=27, color=BLUE, bold=True)
    add_text(slide, "场景的代价随动作变化", 9.88, 3.07, 2.18, 0.34, size=13.6, color=INK, bold=True)
    add_box(slide, 8.52, 4.16, 3.83, 1.55, fill=RED_LIGHT, line=RED)
    add_text(slide, "4 : 1", 8.84, 4.39, 1.38, 0.45, size=27, color=RED, bold=True)
    add_text(slide, "粘附/软地面误判的\n代价方向不对称", 10.14, 4.38, 1.84, 0.73, size=13.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "这一步把“分类准确率”连接到“机器人最终是否付出更大代价”。", y=6.31)

    # 20 — direct recovery
    slide = new_slide(prs)
    add_slide_header(slide, "09 · 从归因到动作", "直接恢复：只在证据充分时释放“后退脱离”动作", 20, source="75 exact-prefix paired trials; current paper Fig. 6")
    sequence = [
        DEMO / "adhesion_peel/review_t01.0s.png",
        DEMO / "adhesion_peel/review_t04.2s.png",
        DEMO / "adhesion_peel/review_t04.8s.png",
    ]
    for i, (p, label) in enumerate(zip(sequence, ["接近异常区", "脚被粘住", "后退脱离"], strict=True)):
        add_image_card(slide, p, 0.72 + i * 2.08, 1.42, 1.88, 1.58, label=label, label_fill=BLACK)
        if i < 2:
            add_text(slide, "→", 2.59 + i * 2.08, 1.98, 0.24, 0.34, size=23, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_picture_contain(slide, assets["recovery"], 7.12, 1.30, 5.25, 3.32)
    add_metric(slide, 0.86, 3.55, 1.88, 1.29, "15 / 15", "释放案例全是粘附", accent=TEAL, fill=TEAL_LIGHT)
    add_metric(slide, 2.94, 3.55, 1.88, 1.29, "20%", "整体动作覆盖率", accent=BLUE, fill=BLUE_LIGHT)
    add_metric(slide, 5.02, 3.55, 1.88, 1.29, "+0.20", "相对等待的成功率", accent=PINK, fill=PINK_LIGHT)
    add_box(slide, 0.84, 5.22, 11.55, 0.90, fill=LIGHT, line=LINE)
    add_text(slide, "安全门规则", 1.10, 5.47, 1.36, 0.34, size=17, color=TEAL, bold=True)
    add_text(slide, "原因是粘附 + 置信度达标 + 5 个模型全部同意 → 后退脱离；否则保持等待。", 2.58, 5.45, 9.45, 0.38, size=15.2, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "这项直接恢复验证聚焦“粘附→后退脱离”；它不是所有异常动作的全覆盖证明。", y=6.35, fill=GRAY_LIGHT, color=MUTED)

    # 21 — evidence chain
    slide = new_slide(prs)
    add_slide_header(slide, "10 · 证据链", "我们现在能够稳健支持的四个结论", 21)
    evidence = [
        ("1", "异常归因是独立问题", "不仅要知道出事，还要知道物理原因。", operator_path("O8", "overview"), BLUE),
        ("2", "两种感知都不可替代", "T2 必须看，T3 必须听身体。", FIG / "fig1_bidirectional_conflict.png", RED),
        ("3", "显式学习冲突更稳", "一套 KiNO 模型覆盖两种证据规则。", FIG / "fig3_structured_router_legacy_boxes_backup.png", TEAL),
        ("4", "正确归因能改善恢复", "动作选择改变代价；粘附恢复已直接验证。", DEMO / "nominal/review_t07.5s.png", PINK),
    ]
    for i, (num, title, body, img, col) in enumerate(evidence):
        x = 0.72 + i * 3.04
        add_box(slide, x, 1.42, 2.82, 4.78, fill=WHITE, line=LINE)
        add_picture_crop(slide, img, x + 0.09, 1.51, 2.64, 1.78, focus_y=0.50)
        add_badge(slide, num, x + 0.20, 1.62, 0.44, fill=col, color=WHITE)
        add_text(slide, title, x + 0.19, 3.57, 2.45, 0.65, size=17.2, color=col, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.25, 4.50, 2.34, 1.00, size=13.1, color=INK, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "四个结论分别由信息极限、跨测试结果、路由审计和物理干预支撑。", y=6.38)

    # 22 — scope and next
    slide = new_slide(prs)
    add_slide_header(slide, "11 · 边界与下一步", "当前证据很完整，但仍有三条明确边界", 22, source="Paper Discussion and independent reconfirmation checkpoint")
    add_image_card(slide, operator_path("O10", "overview"), 0.72, 1.42, 3.00, 2.12, label="难点：驱动力衰减")
    add_image_card(slide, operator_path("O6", "overview"), 3.94, 1.42, 3.00, 2.12, label="难点：外力扰动")
    add_image_card(slide, DEMO / "multiview_peak/03_left_side_low.png", 0.72, 3.84, 6.22, 2.20, label="实机部署仍在计划内")
    add_box(slide, 7.30, 1.42, 5.16, 4.62, fill=LIGHT, line=LINE)
    add_bullets(
        slide,
        [
            ("事件边界由基准提供", "完整系统还要验证何时触发、误报与延迟"),
            ("直接恢复只验证了粘附", "其余原因目前有动作后果矩阵，但缺少逐类闭环 actor"),
            ("当前主证据来自仿真", "真实相机延迟、接触噪声和执行器限制仍需实机验证"),
        ],
        7.66,
        1.80,
        4.42,
        2.57,
        size=14.3,
        bullet_color=RED,
    )
    add_box(slide, 7.66, 4.61, 4.44, 1.03, fill=TEAL_LIGHT, line=TEAL)
    add_text(slide, "正在进行：冻结架构后的全新场景确认", 7.91, 4.82, 3.95, 0.34, size=16, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "新场景、新材质、新种子、新参数区间；完成前不作为最终结论。", 7.92, 5.20, 3.94, 0.30, size=10.8, color=MUTED, align=PP_ALIGN.CENTER)
    add_takeaway(slide, "下一步优先级：完成独立确认 → 扩展原因到动作 actor → Go2 实机部署。", y=6.36, fill=TEAL_LIGHT, color=TEAL)

    # 23 — conclusion
    slide = new_slide(prs)
    add_box(slide, 0.72, 0.66, 5.82, 5.98, fill=WHITE, line=LINE)
    add_badge(slide, "TAKE-HOME MESSAGE", 1.06, 1.00, 2.15, fill=TEAL_LIGHT, color=TEAL)
    add_text(slide, "Feel It. See It.\nRecover.", 1.06, 1.60, 4.98, 1.25, size=36, color=INK, bold=True, font=FONT_EN)
    add_text(slide, "先判断为什么失败，\n再选择如何恢复。", 1.10, 3.18, 4.82, 1.02, size=28, color=INK, bold=True)
    add_text(slide, "Kino-Fail 提供可控异常；KiNO 学习每次该信画面、身体还是联合证据；物理干预把归因连接到真实恢复代价。", 1.10, 4.62, 4.82, 1.12, size=15.2, color=MUTED, line_spacing=1.08)
    add_text(slide, "谢谢", 1.08, 6.18, 1.10, 0.34, size=17, color=TEAL, bold=True)
    add_picture_crop(slide, DEMO / "multiview_peak/08_room_wide.png", 6.88, 0.66, 5.74, 3.28, focus_y=0.52)
    add_picture_crop(slide, DEMO / "multiview_peak/03_left_side_low.png", 6.88, 4.10, 5.74, 2.54, focus_y=0.51)
    add_badge(slide, "真实感室内仿真 · 同一异常的远景与近景", 7.18, 5.99, 4.95, fill=TEAL_LIGHT, color=TEAL)
    add_text(slide, "23", 12.10, 7.12, 0.55, 0.18, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)

    # 24 — appendix paper figure 4
    slide = new_slide(prs)
    add_slide_header(slide, "附录 · 论文原图", "跨两套测试的完整方法比较", 24, source="Current paper Fig. 4")
    add_picture_contain(slide, FIG / "fig4_conflict_results.png", 0.72, 1.28, 11.90, 5.46)

    # 25 — appendix paper figure 5
    slide = new_slide(prs)
    add_slide_header(slide, "附录 · 论文原图", "630 次受控动作干预的代价矩阵", 25, source="Current paper Fig. 5")
    add_picture_contain(slide, FIG / "fig5_intervention_consequence.png", 0.72, 1.28, 11.90, 5.46)

    # 26 — appendix paper figure 6
    slide = new_slide(prs)
    add_slide_header(slide, "附录 · 论文原图", "75 组相同前缀的直接恢复比较", 26, source="Current paper Fig. 6")
    add_picture_contain(slide, FIG / "fig6_selective_recovery.png", 0.72, 1.28, 11.90, 5.46)

    prs.core_properties.title = "KiNO whole-work Chinese presentation"
    prs.core_properties.subject = "Feel It, See It, Recover — Cross-modal failure attribution"
    prs.core_properties.author = "KiNO / Kino-Fail"
    prs.core_properties.keywords = "KiNO, Kino-Fail, quadruped, failure attribution, Chinese presentation"
    prs.core_properties.comments = (
        "Chinese receiver-oriented research presentation. Raster imagery only: "
        "Isaac Sim, benchmark sensor records, data charts, and paper PNG figures."
    )
    return prs


def write_notes() -> Path:
    notes = """# KiNO 整体工作中文汇报讲稿

## 1. 封面
这项工作的核心不是再做一个“能走”的导航系统，而是补上导航失败发生以后经常缺失的一步：机器人究竟为什么失败。标题里的 Feel It、See It、Recover 分别对应身体感知、视觉和原因驱动的恢复。

## 2. 机器人卡住了——真正重要的是为什么
同样是速度下降或停滞，背后的物理原因可能完全不同。软地面需要抬高脚步，粘附需要后退脱离，透明障碍需要停住重规划，底盘托底则需要抬身和减速。仅仅发出“异常”警报，不能决定正确动作。

## 3. 错误归因会执行错误动作
这里用软地面和脚底粘附举例。两者表面都可能表现为走不动，但动作正好相反。我们的受控动作试验显示，粘附时误用抬脚动作的代价为 4，而软地面时保守后退的代价为 1，因此错误方向具有明显不对称性。

## 4. 眼睛和身体谁更可信会反转
整篇工作的关键观察是：不存在一个永远最可信的传感器。T2 里身体输入完全相同，只能看场景；T3 里画面完全相同，只能听身体。这就是我们定义的跨模态失败冲突。

## 5. T2：接触前看场景
软地面与粘附在接触前有不同的视觉语义，但模型收到的身体数据逐字节相同。只听身体的模型不是训练得不够好，而是输入里根本没有可区分信息。

## 6. T3：接触后听身体
低摩擦与透明障碍在指定窗口内共享完全相同的五帧画面，但接触后的速度和加速度反应不同。这里视觉没有区分信息，身体反应才是决定性证据。

## 7. 四个问题
工作由四个部分组成：定义一种此前被忽视的问题；构建可控的数据集；提出显式学习证据选择的方法；最后证明原因判断会改变动作和物理后果。

## 8. Kino-Fail 基准
Kino-Fail 使用 Isaac Sim 的 Unitree Go2，包含 11 个异常算子。冻结数据有 988 个反事实对、1,976 个物理 episode 和 5,928 条同步视觉记录。每个 episode 还带有多个独立外观，覆盖 18 个材质家族。

## 9. 11 个算子（1/2）
前六类覆盖低摩擦、软地面、塌陷、粘附、负载变化和侧向外力。每个卡片的大图是机器人所在场景，小图是同一时刻、同一物理状态的低接触解释视角。

## 10. 11 个算子（2/2）
后五类覆盖外观与物理错配、透明障碍、底盘托底、驱动力衰减和身体感知偏差。它们强调真实故障不一定能从一张图直接看见，身体传感本身也可能出错。

## 11. 外观与物理分开随机化
我们刻意避免把某种物理标签和某张贴图绑定。同一种物理可以换不同材质；同一种外观也可以承载不同物理。物理标签来自仿真算子，纹理只改变画面。

## 12. 模型输入
模型只接收五帧前视 RGB 和一段 19 路身体数据。身体数据再被转成倾斜、滑移、支撑、速度和用力程度等易解释信号。场景编号、算子真值和材质编号都不作为输入。

## 13. 反事实配对
每个案例从干净重置开始，正常侧和异常侧共享场景、起点、路线和相机，只切换一个物理算子。这样能把方法差异与导航随机性分开。

## 14. KiNO 方法
KiNO 有视觉专家、身体专家和联合专家。路由器不读取真值，只观察各专家的置信度、分歧和不确定性，决定这次该信哪一路。最后安全门只有在验证过的原因且多模型一致时才释放动作。

## 15. 两种冲突走不同路线
在 T2，身体输入相同，路由器应选择视觉；在 T3，画面相同，路由器应选择身体。这展示了方法不是固定偏向某种传感器。

## 16. 一套模型，两种测试
横轴是 11 类全量测试，纵轴是 4 类冲突测试。KiNO 不是每个单项的第一名，但它在两种规则反转时保持在右上角：等权平均 97.5%，较差一套测试也有 96.9%。

## 17. 路由是否真的选对
只看视觉专家与身体专家意见不一致的决策行，KiNO 在 T2 有 93.0% 的证据路线正确率，在 T3 为 100%。在同等覆盖率下，置信度还可以把低把握案例留给回退动作。

## 18. 换场景和材质
九场景留一验证平均为 93.8%，最弱测试材质仍有 95.2%。更严格地要求五个训练种子对同一整对样本全部答对时，全量测试通过 295/329 对，冲突测试通过 27/30 案例。

## 19. 从标签到物理后果
我们冻结场景并强制执行七类动作，在九个场景上完成 630 次物理干预。九个场景的最终代价都随动作变化，说明原因标签不是装饰性解释，而是会改变机器人结果的控制接口。

## 20. 直接恢复
在 75 组相同决策前缀的试验中，安全门只释放 15 个案例，并且全部是粘附。相对原地等待，平均代价下降 5.93，成功率提升 0.20；相对后期平均和直接拼接也有更低代价。

## 21. 四个结论
第一，异常归因应作为独立导航问题；第二，视觉和身体感知互不可替代；第三，显式学习冲突路由能在两种证据规则之间保持稳定；第四，正确归因会改变恢复动作和物理代价。

## 22. 边界与下一步
当前实验仍由基准提供事件边界，直接闭环恢复只验证了粘附，主要证据仍来自仿真。下一步依次完成冻结后的全新场景确认、扩展更多原因到动作 actor，并部署到真实 Go2。

## 23. 总结
一句话总结：先判断为什么失败，再决定如何恢复。Kino-Fail 提供可控且更真实的异常，KiNO 学习每次该信画面、身体还是联合证据，物理干预把判断连接到真实代价。

## 24–26. 附录
三页保留论文当前使用的完整结果图，供导师进一步追问跨测试比较、630 次动作干预和 75 组直接恢复的细节。
"""
    path = OUT_DIR / "KiNO_work_chinese_presentation_notes.md"
    path.write_text(notes, encoding="utf-8")
    return path


def main() -> None:
    data = verify_sources()
    assets = prepare_assets(data)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    deck = build_deck(data, assets)
    deck.save(OUT_PPTX)
    notes = write_notes()
    print(f"Wrote {OUT_PPTX} ({len(deck.slides)} slides)")
    print(f"Wrote {notes}")


if __name__ == "__main__":
    main()
