#!/usr/bin/env python3
"""Build the editable English KinoVLA academic presentation.

The deck follows the receiver-oriented guidance in
docs/superpowers/EffectiveCommunicationAmaral.pdf: open with a conundrum,
develop one visual storyline, introduce dense evidence gradually, and reuse the
same visual vocabulary in the conclusion.
"""

from __future__ import annotations

import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "KinoVLA_academic_presentation.pptx"

SW = 13.333
SH = 7.5

FONT = "Aptos"
FONT_DISPLAY = "Aptos Display"

NAVY = "17324D"
INK = "1E2933"
MUTED = "627181"
LINE = "D9E1E7"
LIGHT = "F4F7F9"
TEAL = "0B7A75"
TEAL_LIGHT = "DDF2F0"
ORANGE = "D97745"
ORANGE_LIGHT = "FBE8DD"
GREEN = "2E8B6D"
GREEN_LIGHT = "E2F2EA"
RED = "BC4B51"
RED_LIGHT = "F7E5E6"
GOLD = "B78520"
GOLD_LIGHT = "F7EED8"
WHITE = "FFFFFF"
BLACK = "000000"


def _load_json(relpath: str):
    return json.loads((ROOT / relpath).read_text())


def _close(actual: float, expected: float, *, tol: float = 1e-6) -> None:
    if not math.isclose(float(actual), float(expected), rel_tol=0.0, abs_tol=tol):
        raise AssertionError(f"source value drifted: expected {expected}, got {actual}")


def verify_sources() -> None:
    """Fail fast if a headline number no longer matches its machine artifact."""

    a1 = _load_json("outputs/eval/a1/a1_1_certificate.json")
    a1_probe = a1["fresh_c2st"]["matched_O4_O2"]["obs48|T25|logreg"]
    _close(a1_probe["auc"], 0.5000109265172433)
    _close(a1_probe["perm_p"], 0.10963455149501661)
    assert a1["verdict"]["byte_identical"] is True

    a2 = _load_json("outputs/eval/a2/headline.json")
    _close(a2["summary"]["B1"]["balanced"]["all"]["attribution_balanced"], 0.50)
    _close(a2["summary"]["B5-unshaped"]["balanced"]["all"]["attribution_balanced"], 0.80)
    _close(a2["summary"]["B5-conflict"]["balanced"]["all"]["attribution_balanced"], 0.90)
    paired = a2["mcnemar_O4_conflict"]["B5-unshaped__vs__B5-conflict"]["matched_correct_recovery"]
    assert paired["b_lo_right_hi_wrong"] == 0 and paired["c_lo_wrong_hi_right"] == 24
    _close(paired["p_exact_two_sided"], 1.1920928955078125e-7, tol=1e-14)

    a3 = _load_json("outputs/eval/a3/a3_battery.json")
    _close(a3["heatmap"]["B5-conflict-bi"]["T3"]["acc"], 0.917)
    assert a3["t3_detail"]["B5-conflict-bi"]["looks_safe"]["k"] == 128
    assert a3["t3_detail"]["B5-conflict-bi"]["O8"]["k"] == 48
    assert a3["heatmap"]["B5-conflict-bi"]["T5"]["k"] == 0

    a3_v2 = _load_json("outputs/eval/a3_successor_v2/publication_report.json")
    assert a3_v2["status"] == "confirmatory_passed_all_preregistered_gates"
    assert a3_v2["n_snapshots"] == 237 and a3_v2["n_training_seeds"] == 5
    assert a3_v2["acceptance"]["passed"] is True
    for cell in ("T1", "T2", "T3", "T4", "T5"):
        _close(a3_v2["five_seed_summary"]["per_cell"][cell]["mean"], 1.0)
    strict = a3_v2["strict_cluster_certificate"]["overall"]
    assert strict["k"] == 33 and strict["n"] == 33
    _close(strict["exact_ci95"][0], 0.894237)
    appearance_strict = a3_v2["appearance_id_sensitivity"]["strict"]
    assert appearance_strict["k"] == 15 and appearance_strict["n"] == 15

    a4 = _load_json("outputs/eval/a4/a4_results.json")
    assert a4["n_outcomes"] == 630
    _close(a4["M_mean_cost"]["matched_O4_twophase"]["high_step"], 4.0)
    _close(a4["M_mean_cost"]["matched_O2"]["high_step"], 0.0)
    _close(a4["M_mean_cost"]["matched_O2"]["backstep_detour"], 1.0)

    final = _load_json("outputs/eval/a5/a5_6_c4_final_v2.json")
    repl = _load_json("outputs/eval/a5/a5_6_c4_replication.json")
    assert final["n"] == 288 and repl["n"] == 300
    for obj, expected in (
        (final, (0.2083, 1.5417, 1.75, -0.2083)),
        (repl, (0.20, 1.52, 1.72, -0.20)),
    ):
        coverage, ours, safe, delta = expected
        _close(obj["point_estimate"]["coverage"], coverage)
        _close(obj["point_estimate"]["expected_cost"], ours)
        _close(obj["point_estimate"]["baseline_cost"]["safe"], safe)
        _close(obj["point_estimate"]["paired_delta"]["safe"], delta)
        _close(obj["released_interventions"]["attribution_precision"], 1.0)
    assert repl["released_interventions"]["n_appearance_clusters"] == 6
    _close(repl["released_interventions"]["one_sided_exact_sign_p"], 0.015625)

    a7_route = _load_json("outputs/eval/a7/text_schema_existing.json")["arms"]
    _close(a7_route["vision_only"]["attr_ambiguous_greedy"], 0.4583)
    _close(a7_route["latent"]["attr_ambiguous_greedy"], 1.0)
    _close(
        a7_route["text_binned"]["mean_prompt_tokens"] / a7_route["latent"]["mean_prompt_tokens"],
        1.2692556634304206,
    )
    a7_dose = _load_json("outputs/eval/a7/conflict_dose/summary.json")["aggregate"]
    assert a7_dose["best_mean_dose"] == 10
    _close(a7_dose["by_dose"]["10"]["mean_rate"], 0.733)
    a7_filter = _load_json("outputs/eval/a7/cot_filter_summary.json")
    _close(a7_filter["unfiltered_grounding"]["grounded_correct_rate"], 0.604)
    _close(a7_filter["truth_filtered_grounding"]["grounded_correct_rate"], 0.765)
    a7_encoder = _load_json("outputs/eval/a7/encoder_grid/summary.json")
    _close(a7_encoder["aggregate"]["best_theta_mae"]["value"], 0.0279)

    gate_text = (ROOT / "configs/eval/c4_structured_gate_v2.yaml").read_text()
    assert "0.13283753" in gate_text and "0.15150436" in gate_text


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 20,
    color: str = INK,
    bold: bool = False,
    font: str = FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.02,
    italic: bool = False,
    line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.text = text
    # python-pptx represents newline-separated text as multiple paragraphs.
    # Apply the same typography to every paragraph/run so LibreOffice and
    # PowerPoint render multi-line labels consistently.
    for p in tf.paragraphs:
        p.alignment = align
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        if line_spacing is not None:
            p.line_spacing = line_spacing
        for r in p.runs:
            r.font.name = font
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = rgb(color)
    return box


def add_runs(
    slide,
    runs: list[tuple[str, dict]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 20,
    color: str = INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.02,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    for text, opts in runs:
        r = p.add_run()
        r.text = text
        r.font.name = opts.get("font", FONT)
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", False)
        r.font.italic = opts.get("italic", False)
        r.font.color.rgb = rgb(opts.get("color", color))
    return box


def add_shape(
    slide,
    kind,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str = LINE,
    line_width: float = 1.0,
    radius: bool = False,
):
    shape_kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else kind
    shp = slide.shapes.add_shape(shape_kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.color.rgb = rgb(line)
    shp.line.width = Pt(line_width)
    return shp


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str = LINE,
    line_width: float = 1.0,
    radius: bool = True,
):
    return add_shape(
        slide,
        MSO_SHAPE.RECTANGLE,
        x,
        y,
        w,
        h,
        fill=fill,
        line=line,
        line_width=line_width,
        radius=radius,
    )


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = LINE,
    width: float = 1.2,
    dash: str | None = None,
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if dash:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def add_circle(slide, x, y, d, *, fill=WHITE, line=LINE, line_width=1.0):
    return add_shape(
        slide,
        MSO_SHAPE.OVAL,
        x,
        y,
        d,
        d,
        fill=fill,
        line=line,
        line_width=line_width,
    )


def add_pill(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    *,
    fill: str = LIGHT,
    color: str = NAVY,
    size: float = 11,
    line: str | None = None,
):
    add_box(slide, x, y, w, 0.34, fill=fill, line=line or fill, radius=True)
    add_text(
        slide,
        text,
        x + 0.08,
        y + 0.01,
        w - 0.16,
        0.3,
        size=size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_slide_title(slide, kicker: str, title: str, number: int, *, source: str = ""):
    add_text(slide, kicker.upper(), 0.65, 0.28, 4.2, 0.24, size=10, color=TEAL, bold=True)
    add_text(
        slide, title, 0.65, 0.62, 12.0, 0.62, size=28, color=NAVY, bold=True, font=FONT_DISPLAY
    )
    add_line(slide, 0.65, 7.08, 12.68, 7.08, color=LINE, width=0.8)
    if source:
        add_text(slide, source, 0.65, 7.16, 10.7, 0.18, size=8.5, color=MUTED)
    add_text(
        slide, f"{number:02d}", 12.12, 7.14, 0.56, 0.2, size=9, color=MUTED, align=PP_ALIGN.RIGHT
    )


def add_takeaway(slide, text: str, *, y: float = 6.37, color: str = TEAL, fill: str = TEAL_LIGHT):
    add_box(slide, 0.65, y, 12.03, 0.5, fill=fill, line=fill, radius=True)
    add_text(
        slide,
        text,
        0.9,
        y + 0.06,
        11.55,
        0.38,
        size=16,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_chevron(slide, x: float, y: float, w: float = 0.34, h: float = 0.52, color: str = LINE):
    shp = add_shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, fill=color, line=color)
    return shp


def add_eye(slide, cx: float, cy: float, scale: float = 1.0, color: str = TEAL):
    add_shape(
        slide,
        MSO_SHAPE.ARC,
        cx - 0.34 * scale,
        cy - 0.22 * scale,
        0.68 * scale,
        0.44 * scale,
        fill=WHITE,
        line=color,
        line_width=2.2,
    )
    add_circle(slide, cx - 0.08 * scale, cy - 0.08 * scale, 0.16 * scale, fill=color, line=color)


def add_body_signal(slide, x: float, y: float, w: float, h: float, color: str = ORANGE):
    add_line(slide, x, y + h / 2, x + w, y + h / 2, color=LINE, width=0.8)
    pts = [
        (0.00, 0.52),
        (0.12, 0.52),
        (0.20, 0.30),
        (0.29, 0.78),
        (0.39, 0.18),
        (0.50, 0.62),
        (0.63, 0.44),
        (0.75, 0.58),
        (0.87, 0.51),
        (1.00, 0.51),
    ]
    for (xa, ya), (xb, yb) in zip(pts, pts[1:], strict=False):
        add_line(slide, x + xa * w, y + ya * h, x + xb * w, y + yb * h, color=color, width=2.2)


def add_robot(slide, x: float, y: float, scale: float = 1.0, *, color: str = NAVY):
    # A deliberately simple quadruped pictogram; all elements remain editable.
    add_box(
        slide, x + 0.18 * scale, y, 1.18 * scale, 0.46 * scale, fill=color, line=color, radius=True
    )
    add_box(
        slide,
        x + 1.22 * scale,
        y + 0.06 * scale,
        0.42 * scale,
        0.32 * scale,
        fill=color,
        line=color,
        radius=True,
    )
    add_circle(slide, x + 1.48 * scale, y + 0.13 * scale, 0.08 * scale, fill=WHITE, line=WHITE)
    for lx in (0.34, 0.66, 1.02, 1.26):
        add_line(
            slide,
            x + lx * scale,
            y + 0.38 * scale,
            x + (lx - 0.06) * scale,
            y + 0.90 * scale,
            color=color,
            width=3.2,
        )
        add_line(
            slide,
            x + (lx - 0.06) * scale,
            y + 0.90 * scale,
            x + (lx + 0.08) * scale,
            y + 1.15 * scale,
            color=color,
            width=3.2,
        )
    add_line(
        slide,
        x + 0.2 * scale,
        y + 0.13 * scale,
        x - 0.12 * scale,
        y - 0.12 * scale,
        color=color,
        width=2.2,
    )


def add_soft_ground(
    slide, x: float, y: float, w: float, *, fill: str = TEAL_LIGHT, line: str = TEAL
):
    add_box(slide, x, y, w, 0.48, fill=fill, line=fill, radius=True)
    for i in range(6):
        cx = x + 0.25 + i * (w - 0.5) / 5
        add_circle(slide, cx, y + 0.12 + (i % 2) * 0.08, 0.09, fill=line, line=line)


def add_sticky_ground(
    slide, x: float, y: float, w: float, *, fill: str = ORANGE_LIGHT, line: str = ORANGE
):
    add_box(slide, x, y, w, 0.48, fill=fill, line=fill, radius=True)
    for i in range(5):
        cx = x + 0.35 + i * (w - 0.7) / 4
        add_line(slide, cx, y + 0.08, cx + 0.10, y + 0.35, color=line, width=1.6)


def add_metric_card(slide, x, y, w, h, value, label, *, accent=TEAL, detail=""):
    add_box(slide, x, y, w, h, fill=WHITE, line=LINE, line_width=1.1, radius=True)
    add_box(slide, x, y, 0.08, h, fill=accent, line=accent, radius=False)
    add_text(slide, value, x + 0.28, y + 0.18, w - 0.48, 0.55, size=27, color=accent, bold=True)
    add_text(slide, label, x + 0.28, y + 0.78, w - 0.48, 0.42, size=13.5, color=INK, bold=True)
    if detail:
        add_text(slide, detail, x + 0.28, y + 1.17, w - 0.48, h - 1.3, size=10.5, color=MUTED)


def add_bar_chart(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    labels: list[str],
    values: list[float],
    colors: list[str],
    *,
    maximum: float = 1.0,
    fmt: str = ".2f",
    lower_better: bool = False,
    baseline: float | None = None,
):
    chart_y = y + 0.4
    chart_h = h - 1.0
    add_line(slide, x + 0.45, chart_y, x + 0.45, chart_y + chart_h, color=LINE, width=1.0)
    add_line(
        slide, x + 0.45, chart_y + chart_h, x + w - 0.08, chart_y + chart_h, color=LINE, width=1.0
    )
    for tick in (0, maximum / 2, maximum):
        ty = chart_y + chart_h - (tick / maximum) * chart_h
        add_line(slide, x + 0.40, ty, x + w - 0.08, ty, color=LINE, width=0.55)
        add_text(
            slide,
            f"{tick:g}",
            x,
            ty - 0.12,
            0.34,
            0.22,
            size=8.5,
            color=MUTED,
            align=PP_ALIGN.RIGHT,
        )
    if baseline is not None:
        ty = chart_y + chart_h - (baseline / maximum) * chart_h
        add_line(slide, x + 0.45, ty, x + w - 0.08, ty, color=RED, width=1.2, dash="dash")
    n = len(values)
    usable_w = w - 0.75
    slot = usable_w / n
    bar_w = min(0.68, slot * 0.55)
    for i, (lab, val, col) in enumerate(zip(labels, values, colors, strict=True)):
        bh = max(0.01, (val / maximum) * chart_h)
        bx = x + 0.55 + i * slot + (slot - bar_w) / 2
        by = chart_y + chart_h - bh
        add_box(slide, bx, by, bar_w, bh, fill=col, line=col, radius=False)
        add_text(
            slide,
            format(val, fmt),
            bx - 0.15,
            by - 0.34,
            bar_w + 0.3,
            0.3,
            size=11,
            color=col,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            lab,
            bx - 0.36,
            chart_y + chart_h + 0.12,
            bar_w + 0.72,
            0.5,
            size=9.5,
            color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    if lower_better:
        add_text(
            slide,
            "lower is better",
            x + w - 1.35,
            y,
            1.25,
            0.24,
            size=9,
            color=MUTED,
            italic=True,
            align=PP_ALIGN.RIGHT,
        )


def add_section_marker(slide, label: str, x: float, y: float, *, color=TEAL):
    add_circle(slide, x, y, 0.32, fill=color, line=color)
    add_text(
        slide,
        label,
        x + 0.4,
        y - 0.01,
        1.1,
        0.34,
        size=11,
        color=color,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )


def build_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    blank = prs.slide_layouts[6]

    # ------------------------------------------------------------------ 01
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_pill(
        slide, "KinoVLA · RESEARCH UPDATE", 0.7, 0.48, 2.55, fill=TEAL_LIGHT, color=TEAL, size=10
    )
    add_text(
        slide,
        "Feel It, See It,\nRecover",
        0.7,
        1.15,
        6.4,
        1.62,
        size=38,
        color=NAVY,
        bold=True,
        font=FONT_DISPLAY,
    )
    add_text(
        slide,
        "Cross-modal failure attribution for safer\nquadrupedal navigation recovery",
        0.72,
        2.92,
        5.9,
        0.92,
        size=20,
        color=INK,
        line_spacing=1.05,
    )
    add_box(slide, 0.72, 4.22, 5.66, 1.06, fill=LIGHT, line=LIGHT, radius=True)
    add_text(slide, "Same symptom.", 0.98, 4.42, 1.63, 0.36, size=18, color=MUTED, bold=True)
    add_text(slide, "Different cause.", 2.58, 4.42, 1.83, 0.36, size=18, color=ORANGE, bold=True)
    add_text(slide, "Different recovery.", 4.16, 4.42, 2.05, 0.36, size=18, color=TEAL, bold=True)
    add_text(
        slide, "Isaac Sim · Unitree Go2 · A0–A7 evidence", 0.72, 6.7, 5.4, 0.3, size=11, color=MUTED
    )
    # Hero visual
    add_box(slide, 7.18, 0.66, 5.5, 5.95, fill=LIGHT, line=LIGHT, radius=True)
    add_text(
        slide,
        "ONE STUMBLE",
        7.65,
        1.05,
        4.6,
        0.3,
        size=11,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_robot(slide, 9.05, 1.58, 1.25, color=NAVY)
    add_body_signal(slide, 8.46, 3.22, 2.95, 0.62, color=ORANGE)
    add_text(
        slide,
        "effort ↑   speed ↓",
        8.3,
        3.82,
        3.25,
        0.3,
        size=12,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_line(slide, 9.93, 4.23, 8.35, 4.72, color=LINE, width=1.5)
    add_line(slide, 9.93, 4.23, 11.51, 4.72, color=LINE, width=1.5)
    add_soft_ground(slide, 7.68, 4.82, 2.15)
    add_sticky_ground(slide, 10.84, 4.82, 1.35)
    add_text(
        slide,
        "SOFT MUD",
        7.7,
        5.45,
        2.1,
        0.28,
        size=11,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "ADHESION",
        10.68,
        5.45,
        1.65,
        0.28,
        size=11,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "High-step",
        7.88,
        5.86,
        1.7,
        0.32,
        size=16,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Backstep",
        10.68,
        5.86,
        1.65,
        0.32,
        size=16,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "01", 12.12, 7.14, 0.56, 0.2, size=9, color=MUTED, align=PP_ALIGN.RIGHT)

    # ------------------------------------------------------------------ 02
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Why",
        "A stumble is a symptom—not a diagnosis",
        2,
        source="Conceptual example; measured costs from A4",
    )
    add_text(
        slide,
        "The body can report the same slowdown under two opposite mechanics.",
        0.68,
        1.33,
        11.8,
        0.38,
        size=17,
        color=MUTED,
    )
    # shared signal
    add_box(slide, 4.72, 1.88, 3.9, 1.03, fill=ORANGE_LIGHT, line=ORANGE_LIGHT, radius=True)
    add_body_signal(slide, 5.12, 2.08, 3.1, 0.42, color=ORANGE)
    add_text(
        slide,
        "same proprioceptive trace",
        5.11,
        2.5,
        3.14,
        0.26,
        size=11,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_line(slide, 6.67, 2.92, 3.16, 3.52, color=LINE, width=1.4)
    add_line(slide, 6.67, 2.92, 10.17, 3.52, color=LINE, width=1.4)
    # two branches
    add_box(slide, 0.78, 3.48, 5.35, 2.45, fill=TEAL_LIGHT, line=TEAL_LIGHT, radius=True)
    add_soft_ground(slide, 1.2, 3.92, 2.0, fill=WHITE, line=TEAL)
    add_text(slide, "Soft mud", 3.53, 3.74, 2.1, 0.34, size=22, color=TEAL, bold=True)
    add_text(slide, "Feet sink; clearance helps", 3.53, 4.18, 2.1, 0.55, size=14, color=INK)
    add_pill(slide, "HIGH-STEP", 3.5, 5.08, 1.75, fill=TEAL, color=WHITE, size=12)
    add_box(slide, 7.2, 3.48, 5.35, 2.45, fill=ORANGE_LIGHT, line=ORANGE_LIGHT, radius=True)
    add_sticky_ground(slide, 7.64, 3.92, 2.0, fill=WHITE, line=ORANGE)
    add_text(slide, "Adhesion", 9.96, 3.74, 2.1, 0.34, size=22, color=ORANGE, bold=True)
    add_text(slide, "Pushing harder immobilizes", 9.96, 4.18, 2.1, 0.55, size=14, color=INK)
    add_pill(slide, "BACKSTEP", 9.92, 5.08, 1.75, fill=ORANGE, color=WHITE, size=12)
    add_takeaway(
        slide,
        "In this pair, the wrong recovery costs up to 4 physical-cost units.",
        y=6.29,
        color=RED,
        fill=RED_LIGHT,
    )

    # ------------------------------------------------------------------ 03
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Why",
        "One sensor cannot settle every conflict",
        3,
        source="A1–A3 ambiguity constructions",
    )
    add_box(slide, 0.72, 1.42, 5.82, 4.52, fill=WHITE, line=LINE, radius=True)
    add_pill(slide, "VISION IS DECISIVE", 1.04, 1.72, 2.1, fill=TEAL_LIGHT, color=TEAL, size=11)
    add_eye(slide, 1.4, 2.6, 1.0, TEAL)
    add_body_signal(slide, 2.3, 2.35, 3.55, 0.55, ORANGE)
    add_text(
        slide,
        "same body trace",
        2.5,
        2.92,
        3.12,
        0.28,
        size=11,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_soft_ground(slide, 1.1, 3.52, 1.85)
    add_sticky_ground(slide, 3.75, 3.52, 1.85)
    add_text(
        slide, "mud", 1.25, 4.15, 1.55, 0.28, size=13, color=TEAL, bold=True, align=PP_ALIGN.CENTER
    )
    add_text(
        slide,
        "adhesion",
        3.85,
        4.15,
        1.65,
        0.28,
        size=13,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Body-only must guess",
        1.28,
        4.83,
        4.68,
        0.42,
        size=18,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "A1 matched O4/O2",
        1.65,
        5.34,
        3.92,
        0.28,
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 6.79, 1.42, 5.82, 4.52, fill=WHITE, line=LINE, radius=True)
    add_pill(slide, "BODY IS DECISIVE", 7.1, 1.72, 2.12, fill=ORANGE_LIGHT, color=ORANGE, size=11)
    add_eye(slide, 7.52, 2.6, 1.0, TEAL)
    add_text(
        slide,
        "same safe-looking RGB",
        8.3,
        2.42,
        3.6,
        0.35,
        size=13,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 7.25, 3.45, 2.1, 0.68, fill=LIGHT, line=LINE, radius=True)
    add_text(
        slide, "clear path?", 7.37, 3.63, 1.86, 0.28, size=14, color=MUTED, align=PP_ALIGN.CENTER
    )
    add_body_signal(slide, 9.82, 3.5, 2.3, 0.56, ORANGE)
    add_text(
        slide,
        "impact in joints",
        9.83,
        4.15,
        2.28,
        0.26,
        size=12,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Vision-only misses the obstacle",
        7.16,
        4.83,
        4.98,
        0.42,
        size=18,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "A3 invisible-obstacle probe",
        7.61,
        5.34,
        4.08,
        0.28,
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(slide, "Neither vision nor proprioception subsumes the other.", y=6.29)

    # ------------------------------------------------------------------ 04
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Question",
        "Can the robot learn which evidence to trust?",
        4,
        source="Paper-A research question and claim chain",
    )
    add_text(
        slide,
        "Our thesis: failure attribution is a control decision—not merely a caption.",
        0.72,
        1.36,
        11.9,
        0.42,
        size=18,
        color=MUTED,
    )
    cards = [
        ("01", "Resolve", "learn the direction\nof a sensor conflict", TEAL, TEAL_LIGHT),
        ("02", "Attribute", "name the physical cause\nfrom vision + body", NAVY, LIGHT),
        ("03", "Recover", "choose an action by\nmeasured consequence", ORANGE, ORANGE_LIGHT),
        ("04", "Abstain", "fall back unless all\nobservable cues agree", GREEN, GREEN_LIGHT),
    ]
    for i, (num, title, body, col, fill) in enumerate(cards):
        x = 0.73 + i * 3.07
        add_box(slide, x, 2.08, 2.78, 2.92, fill=fill, line=fill, radius=True)
        add_circle(slide, x + 0.22, 2.32, 0.52, fill=col, line=col)
        add_text(
            slide,
            num,
            x + 0.22,
            2.43,
            0.52,
            0.24,
            size=11,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            title,
            x + 0.24,
            3.08,
            2.25,
            0.4,
            size=22,
            color=col,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            body,
            x + 0.25,
            3.75,
            2.24,
            0.84,
            size=14,
            color=INK,
            align=PP_ALIGN.CENTER,
            line_spacing=1.05,
        )
        if i < 3:
            add_chevron(slide, x + 2.87, 3.24, 0.28, 0.52, color=LINE)
    add_takeaway(
        slide,
        "The evidence chain is only as strong as its weakest link.",
        y=5.82,
        color=NAVY,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 05
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Evaluation",
        "We evaluate the decision—not navigation luck",
        5,
        source="A0 shared protocol; A4 interventional evaluation",
    )
    steps = [
        ("1", "Freeze", "record the failure-onset\nobservation once", TEAL),
        ("2", "Decide", "score every method on\nthe same snapshot", NAVY),
        ("3", "Intervene", "force one recovery while\nholding the scene fixed", ORANGE),
        ("4", "Measure", "read physical outcome:\n0 clean → 6 fall", GREEN),
    ]
    for i, (num, title, body, col) in enumerate(steps):
        x = 0.72 + i * 3.1
        add_circle(slide, x + 0.94, 1.68, 0.56, fill=col, line=col)
        add_text(
            slide,
            num,
            x + 0.94,
            1.82,
            0.56,
            0.2,
            size=12,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            title,
            x + 0.25,
            2.42,
            1.94,
            0.36,
            size=20,
            color=col,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, body, x + 0.08, 2.9, 2.28, 0.74, size=13, color=INK, align=PP_ALIGN.CENTER)
        if i < 3:
            add_chevron(slide, x + 2.42, 1.71, 0.38, 0.5, color=LINE)
    add_line(slide, 1.94, 1.96, 10.77, 1.96, color=LINE, width=2.0)
    add_metric_card(
        slide,
        0.73,
        4.1,
        2.82,
        1.6,
        "548",
        "frozen snapshots",
        accent=TEAL,
        detail="A0 taxonomy corpus",
    )
    add_metric_card(
        slide,
        3.76,
        4.1,
        2.82,
        1.6,
        "240",
        "matched conflicts",
        accent=NAVY,
        detail="A2 paired comparison",
    )
    add_metric_card(
        slide,
        6.79,
        4.1,
        2.82,
        1.6,
        "237",
        "frozen-confirm cases",
        accent=ORANGE,
        detail="A3 v2 · 33 inference clusters",
    )
    add_metric_card(
        slide,
        9.82,
        4.1,
        2.82,
        1.6,
        "630",
        "physics episodes",
        accent=GREEN,
        detail="9 scenes × 7 actions × 10 seeds",
    )
    add_takeaway(
        slide,
        "All physical outcomes are from Isaac Sim with the Unitree Go2 asset.",
        y=6.13,
        color=MUTED,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 06
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Approach",
        "See → Feel → Attribute → Act selectively",
        6,
        source="KinoVLA frozen See–Feel–Act stack",
    )
    # Inputs
    add_box(slide, 0.75, 1.62, 2.1, 1.56, fill=TEAL_LIGHT, line=TEAL_LIGHT, radius=True)
    add_eye(slide, 1.8, 2.1, 0.95, TEAL)
    add_text(
        slide, "SEE", 1.23, 2.64, 1.14, 0.28, size=15, color=TEAL, bold=True, align=PP_ALIGN.CENTER
    )
    add_box(slide, 0.75, 3.62, 2.1, 1.56, fill=ORANGE_LIGHT, line=ORANGE_LIGHT, radius=True)
    add_body_signal(slide, 1.03, 3.98, 1.55, 0.48, ORANGE)
    add_text(
        slide,
        "FEEL",
        1.22,
        4.64,
        1.14,
        0.28,
        size=15,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    # Fusion
    add_chevron(slide, 3.15, 2.87, 0.48, 0.62, color=LINE)
    add_box(slide, 3.82, 2.35, 2.32, 2.38, fill=NAVY, line=NAVY, radius=True)
    add_text(
        slide,
        "Cross-modal\nVLA",
        4.13,
        2.88,
        1.7,
        0.8,
        size=23,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_pill(slide, "Kino-Tokens", 4.25, 3.97, 1.48, fill=WHITE, color=NAVY, size=10)
    # Outputs
    add_chevron(slide, 6.46, 2.87, 0.48, 0.62, color=LINE)
    add_box(slide, 7.12, 1.74, 2.28, 1.38, fill=LIGHT, line=LINE, radius=True)
    add_text(
        slide,
        "ATTRIBUTION",
        7.4,
        1.98,
        1.72,
        0.25,
        size=11,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "“soft terrain”",
        7.34,
        2.43,
        1.83,
        0.32,
        size=18,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 7.12, 3.94, 2.28, 1.38, fill=LIGHT, line=LINE, radius=True)
    add_text(
        slide,
        "PROPOSED ACTION",
        7.34,
        4.18,
        1.84,
        0.25,
        size=11,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "High-step",
        7.36,
        4.63,
        1.8,
        0.32,
        size=18,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    # Gate
    add_chevron(slide, 9.68, 2.87, 0.48, 0.62, color=LINE)
    add_box(
        slide, 10.32, 2.13, 2.28, 2.78, fill=GREEN_LIGHT, line=GREEN, line_width=1.4, radius=True
    )
    add_text(
        slide,
        "SEE–FEEL–ACT\nGATE",
        10.65,
        2.48,
        1.62,
        0.75,
        size=18,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_line(slide, 10.76, 3.45, 12.14, 3.45, color=GREEN, width=1.0)
    add_text(
        slide,
        "all cues agree → release\notherwise → Backstep",
        10.56,
        3.73,
        1.8,
        0.68,
        size=12.5,
        color=INK,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(
        slide,
        "The learned model proposes; the observable gate decides whether to trust it.",
        y=5.76,
    )

    # ------------------------------------------------------------------ 07
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "C1 · Necessity",
        "Controlled ambiguity exposes complementary blind spots",
        7,
        source="A1 matched-pair certificate; A3 O8 probe",
    )
    add_box(slide, 0.72, 1.48, 5.86, 4.54, fill=WHITE, line=LINE, radius=True)
    add_pill(
        slide,
        "MATCHED BODY · DIFFERENT CAUSE",
        1.04,
        1.78,
        2.72,
        fill=ORANGE_LIGHT,
        color=ORANGE,
        size=10,
    )
    add_text(
        slide,
        "Proprioception cannot separate\nmud from adhesion",
        1.02,
        2.36,
        4.98,
        0.86,
        size=20,
        color=NAVY,
        bold=True,
    )
    add_metric_card(
        slide,
        1.02,
        3.44,
        2.15,
        1.5,
        "0.500",
        "proprio AUC",
        accent=ORANGE,
        detail="T=25 logistic probe",
    )
    add_metric_card(
        slide,
        3.48,
        3.44,
        2.15,
        1.5,
        "0.110",
        "permutation p",
        accent=MUTED,
        detail="fresh / deep-reset",
    )
    add_text(
        slide,
        "binding input: byte-identical",
        1.05,
        5.27,
        4.97,
        0.3,
        size=12.5,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 6.81, 1.48, 5.8, 4.54, fill=WHITE, line=LINE, radius=True)
    add_pill(
        slide,
        "SAME RGB · DIFFERENT MECHANICS",
        7.13,
        1.78,
        2.72,
        fill=TEAL_LIGHT,
        color=TEAL,
        size=10,
    )
    add_text(
        slide,
        "Vision cannot see an\ninvisible obstacle",
        7.11,
        2.36,
        4.98,
        0.86,
        size=20,
        color=NAVY,
        bold=True,
    )
    add_metric_card(
        slide,
        7.12,
        3.44,
        2.15,
        1.5,
        "0 / 48",
        "vision-only",
        accent=RED,
        detail="correct attributions",
    )
    add_metric_card(
        slide,
        9.58,
        3.44,
        2.15,
        1.5,
        "48 / 48",
        "body-only",
        accent=GREEN,
        detail="correct attributions",
    )
    add_text(
        slide,
        "the decisive modality flips",
        7.15,
        5.27,
        4.97,
        0.3,
        size=12.5,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(slide, "C1: cross-modal sensing is necessary in both directions.", y=6.28)

    # ------------------------------------------------------------------ 08
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "C2 · Learning",
        "Conflict resolution improves when it is explicitly trained",
        8,
        source="A2, n=240 matched conflicts; exact paired test",
    )
    add_text(
        slide,
        "Attribution accuracy on the same held-out conflict snapshots",
        0.72,
        1.36,
        7.15,
        0.32,
        size=14,
        color=MUTED,
    )
    add_bar_chart(
        slide,
        0.72,
        1.72,
        7.05,
        3.98,
        ["Body-only", "Vision-capable\n(no conflict data)", "Conflict-trained"],
        [0.50, 0.80, 0.90],
        [MUTED, TEAL, NAVY],
        maximum=1.0,
        fmt=".2f",
    )
    add_box(slide, 8.22, 1.72, 4.38, 3.98, fill=LIGHT, line=LIGHT, radius=True)
    add_pill(slide, "+24 / 240", 8.66, 2.11, 1.44, fill=GREEN, color=WHITE, size=13)
    add_text(slide, "paired corrections", 10.22, 2.15, 1.72, 0.28, size=13, color=GREEN, bold=True)
    add_text(
        slide, "B5-unshaped → B5-conflict", 8.64, 2.87, 3.52, 0.35, size=16, color=NAVY, bold=True
    )
    add_text(slide, "No regressions on the paired items", 8.64, 3.37, 3.45, 0.3, size=13, color=INK)
    add_text(slide, "Exact McNemar p = 1.19 × 10⁻⁷", 8.64, 3.88, 3.52, 0.3, size=13, color=INK)
    add_text(slide, "3 / 3 training seeds reach 0.90", 8.64, 4.39, 3.52, 0.3, size=13, color=INK)
    add_pill(
        slide,
        "not “always trust the camera”",
        8.65,
        4.98,
        3.12,
        fill=ORANGE_LIGHT,
        color=ORANGE,
        size=11,
    )
    add_takeaway(
        slide, "C2: having vision is not enough; evidence weighting must be learned.", y=6.21
    )

    # ------------------------------------------------------------------ 09
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "C2 · Confirmation",
        "Structured v2 closes the full bidirectional battery",
        9,
        source="A3-v2 frozen confirmation · 237 snapshots · 5 training seeds",
    )
    add_text(
        slide,
        "Attribution accuracy on 15 new appearances; no confirmation item was used for fitting.",
        0.73,
        1.34,
        11.8,
        0.34,
        size=15,
        color=MUTED,
    )
    for i, cell in enumerate(("T1", "T2", "T3", "T4", "T5")):
        x = 0.74 + i * 2.43
        add_box(slide, x, 1.82, 2.14, 1.25, fill=GREEN_LIGHT, line=GREEN_LIGHT, radius=True)
        add_text(
            slide,
            cell,
            x + 0.2,
            2.04,
            1.74,
            0.28,
            size=14,
            color=GREEN,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            "1.000",
            x + 0.2,
            2.48,
            1.74,
            0.4,
            size=24,
            color=GREEN,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    add_box(slide, 0.75, 3.28, 11.82, 0.62, fill=LIGHT, line=LIGHT, radius=True)
    add_text(slide, "T3 subcells", 1.03, 3.47, 1.35, 0.25, size=12, color=NAVY, bold=True)
    for i, label in enumerate(("looks-safe  1.000", "O8  1.000", "reverse  1.000")):
        add_pill(slide, label, 2.58 + i * 3.0, 3.42, 2.45, fill=GREEN_LIGHT, color=GREEN, size=11)

    add_box(slide, 0.75, 4.15, 7.45, 1.78, fill=WHITE, line=LINE, radius=True)
    add_text(
        slide,
        "Paired case × appearance evidence",
        1.02,
        4.39,
        4.0,
        0.28,
        size=14,
        color=NAVY,
        bold=True,
    )
    comparisons = [
        ("vs body-only", "+0.545", "CI [0.364, 0.727]", "p = 7.63×10⁻⁶"),
        ("vs closed-set fusion", "+0.864", "CI [0.742, 0.962]", "p = 3.73×10⁻⁹"),
        ("vs VLA conflict baseline", "+0.818", "CI [0.667, 0.939]", "p = 1.49×10⁻⁸"),
    ]
    for i, (name, delta, ci, p_value) in enumerate(comparisons):
        y = 4.84 + i * 0.34
        add_text(slide, name, 1.02, y, 2.35, 0.25, size=11.5, color=INK)
        add_text(
            slide,
            delta,
            3.48,
            y,
            0.76,
            0.25,
            size=11.5,
            color=GREEN,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, ci, 4.36, y, 1.75, 0.25, size=11, color=MUTED, align=PP_ALIGN.CENTER)
        add_text(slide, p_value, 6.12, y, 1.72, 0.25, size=10.5, color=MUTED, align=PP_ALIGN.RIGHT)

    add_box(slide, 8.43, 4.15, 4.14, 1.78, fill=GREEN_LIGHT, line=GREEN_LIGHT, radius=True)
    add_text(
        slide,
        "Robustness certificate",
        8.75,
        4.39,
        3.5,
        0.28,
        size=14,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    certificates = [
        "5 / 5 training seeds pass every gate",
        "33 / 33 strict clusters · CI [0.894, 1.000]",
        "15 / 15 appearance IDs · CI [0.782, 1.000]",
        "237 / 237 cause + admissible action",
    ]
    for i, line in enumerate(certificates):
        add_text(
            slide,
            line,
            8.76,
            4.83 + i * 0.28,
            3.48,
            0.23,
            size=10.5,
            color=INK,
            bold=i == 0,
            align=PP_ALIGN.CENTER,
        )
    add_takeaway(
        slide,
        "Structured v2 meets every pre-registered gate and significantly beats "
        "same-battery baselines.",
        y=6.25,
        color=NAVY,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 10
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "C3 · Consequence",
        "The correct recovery depends on the physical cause",
        10,
        source="A4 forced-action consequence matrix; cost: 0 clean, 6 fall",
    )
    add_text(
        slide,
        "Hold the scene fixed. Change only the recovery action.",
        0.73,
        1.36,
        7.4,
        0.32,
        size=16,
        color=MUTED,
    )
    # matrix headers
    add_text(
        slide,
        "FORCED ACTION",
        4.52,
        1.72,
        5.8,
        0.28,
        size=11,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_pill(slide, "BACKSTEP", 4.63, 2.11, 2.4, fill=ORANGE_LIGHT, color=ORANGE, size=12)
    add_pill(slide, "HIGH-STEP", 7.54, 2.11, 2.4, fill=TEAL_LIGHT, color=TEAL, size=12)
    add_text(
        slide,
        "TRUE CAUSE",
        0.8,
        3.04,
        2.36,
        0.28,
        size=11,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Adhesion",
        0.96,
        3.58,
        2.15,
        0.34,
        size=19,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Soft mud",
        0.96,
        4.92,
        2.15,
        0.34,
        size=19,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    cells = [
        (4.55, 3.06, "1", "slow but safe", GOLD_LIGHT, GOLD),
        (7.46, 3.06, "4", "immobilized", RED_LIGHT, RED),
        (4.55, 4.4, "1", "slow but safe", GOLD_LIGHT, GOLD),
        (7.46, 4.4, "0", "clean recovery", GREEN_LIGHT, GREEN),
    ]
    for x, y, val, lab, fill, col in cells:
        add_box(slide, x, y, 2.55, 1.12, fill=fill, line=WHITE, radius=True)
        add_text(
            slide,
            val,
            x + 0.16,
            y + 0.14,
            0.62,
            0.55,
            size=28,
            color=col,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            lab,
            x + 0.79,
            y + 0.37,
            1.55,
            0.3,
            size=13,
            color=col,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    add_box(slide, 10.45, 2.11, 2.02, 3.4, fill=LIGHT, line=LIGHT, radius=True)
    add_text(
        slide, "630", 10.74, 2.57, 1.44, 0.62, size=31, color=NAVY, bold=True, align=PP_ALIGN.CENTER
    )
    add_text(
        slide,
        "physics episodes",
        10.71,
        3.18,
        1.5,
        0.42,
        size=13,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_line(slide, 10.81, 3.88, 12.12, 3.88, color=LINE, width=1.0)
    add_text(
        slide,
        "9 scenes\n7 actions\n10 seeds",
        10.83,
        4.13,
        1.28,
        0.92,
        size=13,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(
        slide,
        "C3: action choice causally changes physical outcome—and errors are asymmetric.",
        y=6.13,
    )

    # ------------------------------------------------------------------ 11
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "C3 · Decision",
        "Asymmetric harm yields a principled safe default",
        11,
        source="A4 T2 belief-weighted consequence analysis",
    )
    add_text(
        slide,
        "Let p be the probability that the patch is adhesive rather than mud.",
        0.73,
        1.36,
        9.6,
        0.32,
        size=16,
        color=MUTED,
    )
    # chart
    x, y, w, h = 0.95, 2.0, 7.2, 3.65
    add_line(slide, x + 0.55, y + h - 0.45, x + w - 0.2, y + h - 0.45, color=INK, width=1.2)
    add_line(slide, x + 0.55, y + 0.2, x + 0.55, y + h - 0.45, color=INK, width=1.2)
    for t in range(5):
        ty = y + h - 0.45 - t * (h - 0.65) / 4
        add_line(slide, x + 0.55, ty, x + w - 0.2, ty, color=LINE, width=0.65)
        add_text(
            slide,
            str(t),
            x + 0.12,
            ty - 0.12,
            0.32,
            0.24,
            size=9,
            color=MUTED,
            align=PP_ALIGN.RIGHT,
        )
    for t in (0, 0.25, 0.5, 0.75, 1.0):
        tx = x + 0.55 + t * (w - 0.75)
        add_text(
            slide,
            f"{t:g}",
            tx - 0.18,
            y + h - 0.27,
            0.36,
            0.24,
            size=9,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
    add_text(
        slide,
        "P(adhesion)",
        x + 2.95,
        y + h + 0.06,
        1.65,
        0.28,
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "expected cost", x - 0.14, y + 0.2, 1.1, 0.28, size=10, color=MUTED)
    # Backstep y=1, high-step y=4p
    plot_x0, plot_x1 = x + 0.55, x + w - 0.2
    plot_y0, plot_y1 = y + h - 0.45, y + 0.2
    y_cost1 = plot_y0 - (1 / 4) * (plot_y0 - plot_y1)
    add_line(slide, plot_x0, y_cost1, plot_x1, y_cost1, color=ORANGE, width=3.0)
    add_line(slide, plot_x0, plot_y0, plot_x1, plot_y1, color=TEAL, width=3.0)
    x_cross = plot_x0 + 0.25 * (plot_x1 - plot_x0)
    add_line(slide, x_cross, plot_y0, x_cross, plot_y1, color=RED, width=1.3, dash="dash")
    add_circle(slide, x_cross - 0.065, y_cost1 - 0.065, 0.13, fill=RED, line=RED)
    add_pill(slide, "p* = 0.25", x_cross - 0.67, y + 0.42, 1.35, fill=RED_LIGHT, color=RED, size=12)
    add_pill(
        slide,
        "Backstep = 1",
        x + 4.93,
        y_cost1 - 0.48,
        1.44,
        fill=ORANGE_LIGHT,
        color=ORANGE,
        size=10,
    )
    add_pill(slide, "High-step = 4p", x + 4.93, y + 0.7, 1.48, fill=TEAL_LIGHT, color=TEAL, size=10)
    # right decision card
    add_box(slide, 8.72, 2.03, 3.8, 3.62, fill=LIGHT, line=LIGHT, radius=True)
    add_text(
        slide,
        "If adhesion is even\nmoderately plausible…",
        9.1,
        2.44,
        3.04,
        0.82,
        size=20,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_pill(slide, "p > 0.25", 9.85, 3.47, 1.56, fill=RED, color=WHITE, size=13)
    add_text(
        slide,
        "prefer Backstep",
        9.04,
        4.12,
        3.12,
        0.38,
        size=20,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "for this mud / adhesion pair",
        9.25,
        4.76,
        2.72,
        0.34,
        size=12,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(
        slide, "A safe default is a consequence-aware decision—not a confidence threshold.", y=6.17
    )

    # ------------------------------------------------------------------ 12
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "C4 · Selective recovery",
        "Release only when See, Feel, and Act agree",
        12,
        source="A5.6 frozen structured gate; deployment-observable inputs only",
    )
    add_text(
        slide,
        "Four observable agreements form a frozen, auditable release rule.",
        0.72,
        1.36,
        11.9,
        0.38,
        size=17,
        color=MUTED,
    )
    gate_items = [
        ("SEE", "RGB resembles calibrated\nsoft-terrain materials", TEAL, TEAL_LIGHT),
        ("FEEL", "tracking-error peak confirms\na physical anomaly", ORANGE, ORANGE_LIGHT),
        ("NAME", "attribution is\n“compliant terrain”", NAVY, LIGHT),
        ("ACT", "proposed recovery is\nHigh-step", GREEN, GREEN_LIGHT),
    ]
    for i, (head, body, col, fill) in enumerate(gate_items):
        x = 0.74 + i * 3.03
        add_box(slide, x, 2.1, 2.7, 1.82, fill=fill, line=fill, radius=True)
        add_circle(slide, x + 0.19, 2.32, 0.43, fill=col, line=col)
        add_text(
            slide,
            "✓",
            x + 0.19,
            2.4,
            0.43,
            0.21,
            size=12,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, head, x + 0.75, 2.31, 1.58, 0.3, size=14, color=col, bold=True)
        add_text(slide, body, x + 0.27, 2.88, 2.17, 0.7, size=13, color=INK, align=PP_ALIGN.CENTER)
        if i < 3:
            add_text(
                slide,
                "+",
                x + 2.77,
                2.77,
                0.42,
                0.42,
                size=23,
                color=MUTED,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
    add_line(slide, 6.66, 4.08, 6.66, 4.44, color=GREEN, width=2.0)
    add_shape(slide, MSO_SHAPE.DOWN_ARROW, 6.45, 4.31, 0.42, 0.56, fill=GREEN, line=GREEN)
    add_box(slide, 3.93, 4.84, 2.38, 0.92, fill=GREEN, line=GREEN, radius=True)
    add_text(
        slide,
        "ALL TRUE\nrelease High-step",
        4.17,
        5.02,
        1.9,
        0.54,
        size=14,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 7.02, 4.84, 2.38, 0.92, fill=MUTED, line=MUTED, radius=True)
    add_text(
        slide,
        "ANY FALSE\nuse Backstep",
        7.26,
        5.02,
        1.9,
        0.54,
        size=14,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_pill(
        slide,
        "forbidden: truth · scenario ID · appearance ID · θ · split · cost",
        3.58,
        5.98,
        6.18,
        fill=RED_LIGHT,
        color=RED,
        size=10,
    )
    add_takeaway(
        slide,
        "The VLA, projector, training data, and A1–A4 evidence remain frozen.",
        y=6.39,
        color=NAVY,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 13
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "C4 · Result",
        "The frozen gate beats always-safe on two final evaluations",
        13,
        source="A5.6 final-v2 and confirmatory replication; paired two-stage bootstrap",
    )
    add_text(
        slide,
        "Expected physical cost (lower is better)",
        0.74,
        1.38,
        5.2,
        0.3,
        size=14,
        color=MUTED,
    )
    # two-group chart, custom for comparison
    add_bar_chart(
        slide,
        0.74,
        1.78,
        7.55,
        4.15,
        [
            "Final-v2\nours",
            "Final-v2\nalways-safe",
            "Replication\nours",
            "Replication\nalways-safe",
        ],
        [1.5417, 1.75, 1.52, 1.72],
        [TEAL, MUTED, TEAL, MUTED],
        maximum=2.0,
        fmt=".2f",
        lower_better=True,
    )
    add_box(slide, 8.67, 1.78, 3.9, 4.15, fill=LIGHT, line=LIGHT, radius=True)
    add_text(slide, "FINAL-v2", 9.0, 2.12, 1.32, 0.25, size=11, color=MUTED, bold=True)
    add_text(slide, "Δ = −0.208", 9.0, 2.47, 2.55, 0.4, size=23, color=TEAL, bold=True)
    add_text(slide, "95% CI [−0.208, −0.208]", 9.0, 2.92, 2.95, 0.3, size=12, color=INK)
    add_line(slide, 9.0, 3.38, 12.12, 3.38, color=LINE, width=1.0)
    add_text(slide, "REPLICATION", 9.0, 3.7, 1.6, 0.25, size=11, color=MUTED, bold=True)
    add_text(slide, "Δ = −0.200", 9.0, 4.05, 2.55, 0.4, size=23, color=TEAL, bold=True)
    add_text(slide, "95% CI [−0.20, −0.20]", 9.0, 4.5, 2.95, 0.3, size=12, color=INK)
    add_line(slide, 9.0, 4.96, 12.12, 4.96, color=LINE, width=1.0)
    add_text(slide, "release precision 1.00", 9.0, 5.22, 2.95, 0.3, size=13, color=GREEN, bold=True)
    add_takeaway(
        slide, "C4: in-support selective recovery is safer than always using the fallback.", y=6.21
    )

    # ------------------------------------------------------------------ 14
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "C4 · Confirmation",
        "The positive result survives a post-freeze appearance replication",
        14,
        source="A5.6 six-cluster confirmatory corpus; unchanged gate hash",
    )
    # timeline
    add_line(slide, 1.28, 2.19, 11.93, 2.19, color=LINE, width=2.4)
    stages = [
        (1.18, "Develop", "select observable\nconditions", MUTED),
        (4.31, "Freeze", "gate hash\n1d538a…9287bbd", NAVY),
        (7.47, "Final-v2", "288 snapshots\n2 release clusters", TEAL),
        (10.57, "Replicate", "300 snapshots\n24 new appearance IDs", GREEN),
    ]
    for x, head, body, col in stages:
        add_circle(slide, x, 1.94, 0.5, fill=col, line=col)
        add_text(
            slide,
            head,
            x - 0.45,
            2.62,
            1.4,
            0.3,
            size=15,
            color=col,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, body, x - 0.66, 3.12, 1.82, 0.65, size=12, color=INK, align=PP_ALIGN.CENTER)
    add_box(slide, 1.1, 4.25, 11.15, 1.38, fill=GREEN_LIGHT, line=GREEN_LIGHT, radius=True)
    add_text(
        slide,
        "6 / 6",
        1.52,
        4.56,
        1.25,
        0.52,
        size=29,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "released mud appearance clusters improved by exactly 1.0 cost unit",
        3.05,
        4.55,
        7.88,
        0.38,
        size=17,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_pill(
        slide,
        "one-sided exact sign test p = 0.015625",
        4.12,
        5.08,
        4.48,
        fill=WHITE,
        color=GREEN,
        size=11,
        line=GREEN,
    )
    add_takeaway(
        slide, "Independent appearances, unchanged gate, no released-cluster harm.", y=6.21
    )

    # ------------------------------------------------------------------ 15
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "A7 · What made it work",
        "The load-bearing choices are measurable",
        15,
        source="A7 controlled route, dose, rationale, and encoder ablations",
    )
    # Panel 1 route
    add_box(slide, 0.72, 1.49, 3.84, 4.6, fill=WHITE, line=LINE, radius=True)
    add_pill(slide, "BODY ROUTE", 1.02, 1.79, 1.3, fill=ORANGE_LIGHT, color=ORANGE, size=10)
    add_text(
        slide, "0.458", 1.0, 2.48, 1.35, 0.55, size=28, color=RED, bold=True, align=PP_ALIGN.CENTER
    )
    add_text(
        slide, "vision-only", 1.05, 3.03, 1.26, 0.28, size=12, color=MUTED, align=PP_ALIGN.CENTER
    )
    add_text(
        slide, "→", 2.36, 2.58, 0.46, 0.38, size=24, color=LINE, bold=True, align=PP_ALIGN.CENTER
    )
    add_text(
        slide,
        "1.000",
        2.79,
        2.48,
        1.35,
        0.55,
        size=28,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "latent", 2.84, 3.03, 1.26, 0.28, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    add_line(slide, 1.1, 3.64, 4.15, 3.64, color=LINE, width=1.0)
    add_text(
        slide,
        "same VLM harness",
        1.23,
        3.91,
        2.78,
        0.28,
        size=13,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "binned text uses 1.27× as many\ntokens; latent adds a θ head",
        1.1,
        4.49,
        3.06,
        0.72,
        size=13,
        color=ORANGE,
        align=PP_ALIGN.CENTER,
    )
    # Panel 2 dose
    add_box(slide, 4.75, 1.49, 3.84, 4.6, fill=WHITE, line=LINE, radius=True)
    add_pill(slide, "CONFLICT DOSE", 5.05, 1.79, 1.55, fill=TEAL_LIGHT, color=TEAL, size=10)
    vals = [0.533, 0.333, 0.733, 0.667, 0.667]
    labs = ["0", "5", "10", "20", "40"]
    px0, py0, pw, ph = 5.17, 2.52, 2.98, 1.73
    add_line(slide, px0, py0 + ph, px0 + pw, py0 + ph, color=LINE, width=1.0)
    for i in range(4):
        add_line(slide, px0, py0 + i * ph / 3, px0 + pw, py0 + i * ph / 3, color=LINE, width=0.55)
    pts = []
    for i, (lab, val) in enumerate(zip(labs, vals, strict=True)):
        cx = px0 + i * pw / 4
        cy = py0 + ph - val * ph
        pts.append((cx, cy))
        add_circle(
            slide,
            cx - 0.07,
            cy - 0.07,
            0.14,
            fill=TEAL if lab != "10" else GREEN,
            line=TEAL if lab != "10" else GREEN,
        )
        add_text(
            slide,
            lab,
            cx - 0.2,
            py0 + ph + 0.08,
            0.4,
            0.24,
            size=9,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
    for a, b in zip(pts, pts[1:], strict=False):
        add_line(slide, *a, *b, color=TEAL, width=2.0)
    add_pill(
        slide,
        "best mean = 0.733 at dose 10",
        5.16,
        4.67,
        2.95,
        fill=GREEN_LIGHT,
        color=GREEN,
        size=10,
    )
    add_text(
        slide,
        "3 seeds · non-monotone",
        5.27,
        5.23,
        2.72,
        0.28,
        size=12,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    # Panel 3 rigor
    add_box(slide, 8.78, 1.49, 3.84, 4.6, fill=WHITE, line=LINE, radius=True)
    add_pill(slide, "GROUNDING", 9.08, 1.79, 1.28, fill=GREEN_LIGHT, color=GREEN, size=10)
    add_text(
        slide,
        "0.604",
        9.05,
        2.48,
        1.35,
        0.55,
        size=28,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "raw rationales",
        9.05,
        3.03,
        1.36,
        0.28,
        size=11.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide, "→", 10.43, 2.58, 0.46, 0.38, size=24, color=LINE, bold=True, align=PP_ALIGN.CENTER
    )
    add_text(
        slide,
        "0.765",
        10.88,
        2.48,
        1.35,
        0.55,
        size=28,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "truth-filtered",
        10.87,
        3.03,
        1.37,
        0.28,
        size=11.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_line(slide, 9.14, 3.64, 12.2, 3.64, color=LINE, width=1.0)
    add_text(slide, "best θ MAE", 9.16, 4.03, 1.55, 0.28, size=13, color=INK, bold=True)
    add_text(
        slide,
        "0.0279",
        10.66,
        3.91,
        1.48,
        0.5,
        size=25,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    add_text(
        slide,
        "privileged-distillation encoder",
        9.15,
        4.61,
        2.96,
        0.3,
        size=11.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(
        slide,
        "Proprioceptive evidence, conflict supervision, and grounded filtering are load-bearing.",
        y=6.29,
        color=NAVY,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 16
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Scope",
        "What the evidence supports—and what it does not",
        16,
        source="A0–A7 evidence audit; negative results retained as boundaries",
    )
    add_box(slide, 0.73, 1.52, 5.83, 4.67, fill=GREEN_LIGHT, line=GREEN_LIGHT, radius=True)
    add_pill(slide, "SUPPORTED", 1.06, 1.85, 1.36, fill=GREEN, color=WHITE, size=11)
    supported = [
        "Complementary single-modality blind spots",
        "Learned bidirectional conflict resolution",
        "Frozen all-cell structured confirmation",
        "Causal action–consequence asymmetry",
        "In-support structured selective recovery",
    ]
    for i, txt in enumerate(supported):
        y = 2.55 + i * 0.65
        add_circle(slide, 1.08, y + 0.02, 0.25, fill=GREEN, line=GREEN)
        add_text(
            slide,
            "✓",
            1.08,
            y + 0.06,
            0.25,
            0.13,
            size=8.5,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, txt, 1.52, y - 0.02, 4.43, 0.32, size=14, color=INK)
    add_box(slide, 6.79, 1.52, 5.82, 4.67, fill=RED_LIGHT, line=RED_LIGHT, radius=True)
    add_pill(slide, "NOT CLAIMED", 7.13, 1.85, 1.55, fill=RED, color=WHITE, size=11)
    limits = [
        "Physical-robot deployment",
        "Open-world material recognition",
        "Unseen-operator semantic naming",
        "A universal scalar abstention score",
        "A learned intervention boundary",
        "End-to-end VLA all-cell dominance",
    ]
    for i, txt in enumerate(limits):
        y = 2.5 + i * 0.56
        add_circle(slide, 7.13, y + 0.03, 0.25, fill=RED, line=RED)
        add_text(
            slide,
            "×",
            7.13,
            y + 0.055,
            0.25,
            0.15,
            size=10,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, txt, 7.57, y - 0.01, 4.45, 0.32, size=13.5, color=INK)
    add_takeaway(
        slide,
        "The positive C4 claim is narrow by design: calibrated material support, "
        "not open-world OOD.",
        y=6.35,
        color=NAVY,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 17
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Takeaways",
        "Feel it. See it. Recover selectively.",
        17,
        source="KinoVLA Paper-A · discussion",
    )
    # Reuse visual language from opening.
    add_robot(slide, 0.96, 1.72, 1.1, color=NAVY)
    add_soft_ground(slide, 0.78, 3.12, 2.05)
    add_sticky_ground(slide, 0.78, 3.79, 2.05)
    takeaways = [
        (
            "1",
            "Ambiguity is real",
            "The same symptom can hide opposite mechanics.",
            TEAL,
            TEAL_LIGHT,
        ),
        (
            "2",
            "Conflict resolution is robust",
            "Structured v2 passes every battery cell across five training seeds.",
            NAVY,
            LIGHT,
        ),
        (
            "3",
            "Recovery must be selective",
            "A frozen observable gate improves cost on two final sets.",
            GREEN,
            GREEN_LIGHT,
        ),
    ]
    for i, (num, head, body, col, fill) in enumerate(takeaways):
        y = 1.48 + i * 1.53
        add_box(slide, 3.43, y, 8.95, 1.15, fill=fill, line=fill, radius=True)
        add_circle(slide, 3.7, y + 0.26, 0.58, fill=col, line=col)
        add_text(
            slide,
            num,
            3.7,
            y + 0.41,
            0.58,
            0.2,
            size=13,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(slide, head, 4.57, y + 0.13, 4.35, 0.52, size=17, color=col, bold=True)
        add_text(slide, body, 4.57, y + 0.7, 7.17, 0.3, size=13.5, color=INK)
    add_pill(slide, "DISCUSSION", 0.97, 5.25, 1.64, fill=NAVY, color=WHITE, size=12)
    add_text(
        slide,
        "Where should we expand next: broader support, new mechanics, or hardware?",
        3.45,
        6.24,
        8.96,
        0.42,
        size=18,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    # =============================================================== APPENDIX
    # ------------------------------------------------------------------ 18
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Appendix",
        "A0–A7 evidence map",
        18,
        source="Machine artifacts and frozen reducers; no external-dataset evidence",
    )
    exps = [
        ("A0", "Protocol", "548 snapshots\ndeterminism + corpus", MUTED),
        ("A1", "Necessity", "matched-body\nconstruct validity", ORANGE),
        ("A2", "Vision-true", "240 paired\nconflicts", TEAL),
        ("A3", "Conflict", "237 frozen confirm\n33 strict clusters", NAVY),
        ("A4", "Consequence", "630 physics\nepisodes", RED),
        ("A5", "Selective", "288 + 300\nfinal snapshots", GREEN),
        ("A6", "Boundary", "base bracket;\nno learned flip", GOLD),
        ("A7", "Ablations", "route, dose,\ngrounding, uncertainty", TEAL),
    ]
    for i, (exp, head, body, col) in enumerate(exps):
        row, c = divmod(i, 4)
        x = 0.73 + c * 3.06
        y = 1.52 + row * 2.38
        add_box(slide, x, y, 2.76, 1.96, fill=WHITE, line=LINE, radius=True)
        add_pill(slide, exp, x + 0.22, y + 0.2, 0.64, fill=col, color=WHITE, size=11)
        add_text(slide, head, x + 0.96, y + 0.22, 1.58, 0.3, size=14, color=col, bold=True)
        add_text(
            slide, body, x + 0.24, y + 0.86, 2.28, 0.68, size=13, color=INK, align=PP_ALIGN.CENTER
        )
    add_takeaway(
        slide,
        "Every A0–A7 block has traceable machine-readable artifacts.",
        y=6.31,
        color=NAVY,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 19
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Appendix",
        "Statistical unit and uncertainty by experiment",
        19,
        source="A0–A7 evidence audit",
    )
    headers = ["Experiment", "Primary unit", "Uncertainty / test", "Role"]
    widths = [1.35, 3.35, 4.25, 2.42]
    x0, y0 = 0.78, 1.48
    x = x0
    for head, w in zip(headers, widths, strict=True):
        add_box(slide, x, y0, w, 0.52, fill=NAVY, line=WHITE, radius=False)
        add_text(
            slide,
            head,
            x + 0.08,
            y0 + 0.14,
            w - 0.16,
            0.22,
            size=11,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        x += w
    data = [
        ("A1", "fresh app × lane × seed", "C2ST + permutation", "construct validity"),
        ("A2", "matched pair / appearance", "exact McNemar + 3 seeds", "conflict learning"),
        ("A3-v2", "case×appearance + seed", "2-level bootstrap + sign test", "frozen confirmation"),
        ("A4", "physical episode", "2-stage paired bootstrap", "causal cost"),
        (
            "A5.6",
            "appearance cluster + episode",
            "2-stage paired bootstrap + sign test",
            "selective recovery",
        ),
        ("A7", "training seed / appearance", "mean + range; held-out CIs", "method ablation"),
    ]
    for i, row in enumerate(data):
        y = y0 + 0.52 + i * 0.7
        x = x0
        fill = LIGHT if i % 2 == 0 else WHITE
        for txt, w in zip(row, widths, strict=True):
            add_box(slide, x, y, w, 0.68, fill=fill, line=WHITE, radius=False)
            add_text(
                slide,
                txt,
                x + 0.08,
                y + 0.13,
                w - 0.16,
                0.38,
                size=11.5,
                color=INK,
                bold=(x == x0),
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )
            x += w
    add_takeaway(
        slide,
        "Snapshot counts are descriptive; claims respect appearance, seed, and episode dependence.",
        y=6.24,
        color=NAVY,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 20
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Appendix",
        "A3-v2 paired comparator audit",
        20,
        source="A3-v2 frozen confirmation · cluster-balanced paired inference",
    )
    comparisons = [
        ("BODY-ONLY · B1", "+0.545", "95% CI  [0.364, 0.727]", "p = 7.629 × 10⁻⁶", TEAL),
        ("CLOSED-SET FUSION · B-F", "+0.864", "95% CI  [0.742, 0.962]", "p = 3.725 × 10⁻⁹", NAVY),
        ("VLA CONFLICT BASELINE", "+0.818", "95% CI  [0.667, 0.939]", "p = 1.490 × 10⁻⁸", GREEN),
    ]
    for i, (head, delta, interval, pvalue, col) in enumerate(comparisons):
        x = 0.75 + i * 4.03
        add_box(slide, x, 1.55, 3.72, 2.78, fill=WHITE, line=LINE, radius=True)
        add_pill(slide, head, x + 0.28, 1.83, 3.16, fill=col, color=WHITE, size=9.5)
        add_text(
            slide,
            delta,
            x + 0.3,
            2.55,
            3.12,
            0.52,
            size=27,
            color=col,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            "cluster-balanced accuracy",
            x + 0.38,
            3.13,
            2.96,
            0.26,
            size=11,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            interval,
            x + 0.3,
            3.57,
            3.12,
            0.3,
            size=12.5,
            color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            pvalue,
            x + 0.3,
            3.93,
            3.12,
            0.25,
            size=11.5,
            color=INK,
            align=PP_ALIGN.CENTER,
        )
    add_box(slide, 0.75, 4.72, 11.78, 1.23, fill=GREEN_LIGHT, line=GREEN_LIGHT, radius=True)
    certificates = [
        ("5 / 5", "training seeds pass"),
        ("33 / 33", "strict clusters · CI [0.894, 1.000]"),
        ("15 / 15", "appearance IDs · CI [0.782, 1.000]"),
    ]
    for i, (value, label) in enumerate(certificates):
        x = 1.08 + i * 3.82
        add_text(slide, value, x, 4.96, 1.26, 0.35, size=19, color=GREEN, bold=True)
        add_text(slide, label, x + 1.18, 4.96, 2.28, 0.42, size=11.5, color=INK)
    add_takeaway(
        slide,
        "Every same-battery paired interval excludes zero; exact sign-test evidence is decisive.",
        y=6.31,
        color=GREEN,
        fill=GREEN_LIGHT,
    )

    # ------------------------------------------------------------------ 21
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Appendix",
        "A3-v2 pre-registered gates: all passed",
        21,
        source="outputs/eval/a3_successor_v2/publication_report.json",
    )
    items = [
        ("T1 agree", "1.000", "required ≥ 0.90"),
        ("T2 vision-true", "1.000", "required = 1.00"),
        ("T3 body-true", "1.000", "required > 0.917"),
        ("T4 fine structure", "1.000", "required = 1.00"),
        ("T5 continue", "1.000", "required ≥ 0.90"),
        ("Worst / macro", "1.000", "required ≥ 0.90 / 0.95"),
    ]
    for i, (head, value, body) in enumerate(items):
        row, c = divmod(i, 3)
        x = 0.73 + c * 4.03
        y = 1.52 + row * 2.28
        add_box(slide, x, y, 3.72, 1.9, fill=GREEN_LIGHT, line=GREEN_LIGHT, radius=True)
        add_text(slide, head, x + 0.27, y + 0.23, 2.03, 0.3, size=15, color=GREEN, bold=True)
        add_text(
            slide,
            value,
            x + 2.36,
            y + 0.18,
            1.04,
            0.4,
            size=19,
            color=GREEN,
            bold=True,
            align=PP_ALIGN.RIGHT,
        )
        add_text(
            slide, body, x + 0.27, y + 0.88, 3.14, 0.55, size=13, color=INK, align=PP_ALIGN.CENTER
        )
    add_takeaway(
        slide,
        "Five seeds · method SHA-256 537a5861…1d89c · frozen before confirmation collection.",
        y=6.23,
        color=NAVY,
        fill=LIGHT,
    )

    # ------------------------------------------------------------------ 22
    slide = prs.slides.add_slide(blank)
    add_slide_title(
        slide,
        "Appendix",
        "Frozen See–Feel–Act gate specification",
        22,
        source="configs/eval/c4_structured_gate_v2.yaml; A5.6 final artifacts",
    )
    specs = [
        ("SEE", "material-support distance", "≤ 0.13283753", TEAL, TEAL_LIGHT),
        ("FEEL", "tracking-error peak", "≥ 0.15150436", ORANGE, ORANGE_LIGHT),
        ("NAME", "required attribution", "compliant_terrain", NAVY, LIGHT),
        ("ACT", "required proposal", "high_step", GREEN, GREEN_LIGHT),
    ]
    for i, (head, label, value, col, fill) in enumerate(specs):
        x = 0.73 + i * 3.03
        add_box(slide, x, 1.6, 2.7, 2.18, fill=fill, line=fill, radius=True)
        add_pill(slide, head, x + 0.26, 1.87, 0.92, fill=col, color=WHITE, size=10)
        add_text(
            slide,
            label,
            x + 0.27,
            2.47,
            2.16,
            0.48,
            size=13,
            color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            value,
            x + 0.2,
            3.13,
            2.3,
            0.35,
            size=16,
            color=col,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    add_box(slide, 0.74, 4.17, 5.63, 1.37, fill=GREEN_LIGHT, line=GREEN_LIGHT, radius=True)
    add_text(slide, "ALL FOUR TRUE", 1.02, 4.46, 1.68, 0.28, size=13, color=GREEN, bold=True)
    add_text(
        slide,
        "release High-step",
        2.72,
        4.42,
        3.1,
        0.38,
        size=19,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide, "coverage ≈ 0.20", 2.72, 4.91, 3.1, 0.28, size=12, color=MUTED, align=PP_ALIGN.CENTER
    )
    add_box(slide, 6.61, 4.17, 5.97, 1.37, fill=LIGHT, line=LIGHT, radius=True)
    add_text(slide, "ANY FALSE", 6.91, 4.46, 1.38, 0.28, size=13, color=MUTED, bold=True)
    add_text(
        slide,
        "use backstep_detour",
        8.43,
        4.42,
        3.58,
        0.38,
        size=19,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "gate SHA-256 1d538a…9287bbd",
        8.43,
        4.91,
        3.58,
        0.28,
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_takeaway(
        slide,
        "No truth, scenario ID, appearance ID, θ, split, or cost is read at deployment.",
        y=6.16,
        color=RED,
        fill=RED_LIGHT,
    )

    # Core properties
    prs.core_properties.title = "Feel It, See It, Recover — KinoVLA Academic Presentation"
    prs.core_properties.subject = (
        "Cross-modal failure attribution for safer quadrupedal navigation recovery"
    )
    prs.core_properties.author = "KinoVLA"
    prs.core_properties.keywords = "KinoVLA, quadruped, cross-modal, recovery, attribution, A0-A7"
    prs.core_properties.comments = (
        "Receiver-oriented white-background academic deck following Amaral's "
        "communication guidance."
    )
    return prs


def main() -> None:
    verify_sources()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = build_deck()
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
