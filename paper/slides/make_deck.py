#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Kino-VLA group-meeting deck (6 slides).

Plain Chinese, no jargon-only labels, Claude/Anthropic warm-paper aesthetic.
Single source of truth -> renders BOTH an editable .pptx AND PNG/PDF previews
(the previews let us QA layout + look without PowerPoint/LibreOffice installed).
"""
import math
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

# ----------------------------------------------------------------------------- theme
PAGE = "F5F0E6"      # warm paper background
CARD = "FCFAF5"      # light card
BAND = "EDE3D2"      # deeper warm band
TINT = "F4E2D6"      # coral-tinted band
INK = "23201B"       # near-black warm ink
INK_SOFT = "6E655A"  # secondary text
INK_FAINT = "A99F90"  # captions / page numbers
HAIR = "E2D7C5"      # hairline / divider
CORAL = "CF6F4B"     # Claude clay / terracotta
CORAL_DEEP = "A9542F"
GRAYBAR = "C2B7A6"   # muted bar (baseline comparison)
YELLOW = "E6C24E"
BROWN = "8C5A33"
SOFT1 = "EAD9C2"     # faint decorative spark
SOFT2 = "E7C9B2"

SERIF = "Noto Serif CJK SC"
SANS = "Noto Sans CJK SC"
FAMILY = {"serif": SERIF, "sans": SANS}
FONT_FILES = {
    ("serif", False): ("/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc", 0),
    ("serif", True): ("/usr/share/fonts/opentype/noto/NotoSerifCJK-Bold.ttc", 0),
    ("sans", False): ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", 0),
    ("sans", True): ("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc", 0),
}

PAGE_W, PAGE_H = 13.333, 7.5
ML, MR = 0.85, 0.85
CW = PAGE_W - ML - MR
DPI = 130


def hx(h):
    return RGBColor.from_string(h)


def rgb(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


# ----------------------------------------------------------------------------- data builders
def R(t, size=14.5, color=INK, font="sans", bold=False, italic=False):
    return dict(t=t, size=size, color=color, font=font, bold=bold, italic=italic)


def P(runs, align="l", sa=6.0, sb=0.0, ls=1.14):
    if isinstance(runs, dict):
        runs = [runs]
    return dict(runs=runs, align=align, sa=sa, sb=sb, ls=ls)


def T(x, y, w, h, paras, valign="t"):
    if isinstance(paras, dict):
        paras = [paras]
    return dict(type="text", x=x, y=y, w=w, h=h, paras=paras, valign=valign)


def Rt(x, y, w, h, fill=None, line=None, lw=1.0, round=0.0):
    return dict(type="rect", x=x, y=y, w=w, h=h, fill=fill, line=line, lw=lw, round=round)


def Ln(x1, y1, x2, y2, color=CORAL, w=1.5):
    return dict(type="line", x1=x1, y1=y1, x2=x2, y2=y2, color=color, w=w)


def spark(cx, cy, r, color=CORAL, n=8, w=2.0):
    out = []
    for i in range(n):
        a = math.pi * 2 * i / n
        out.append(Ln(cx, cy, cx + r * math.cos(a), cy + r * math.sin(a), color=color, w=w))
    return out


CIRC = ["①", "②", "③", "④"]


# ----------------------------------------------------------------------------- shared chrome
def header(page, eyebrow, title, title_size=27):
    E = []
    E += spark(ML + 0.08, 0.67, 0.12, CORAL, 8, 1.7)
    E.append(T(ML + 0.32, 0.51, 9.0, 0.4, P(R(eyebrow, 13, CORAL, "sans", True), sa=0)))
    E.append(T(ML, 0.84, CW, 0.95, P(R(title, title_size, INK, "serif", True), ls=1.04, sa=0)))
    E.append(Ln(ML + 0.02, 1.73, ML + 1.7, 1.73, CORAL, 2.4))
    E.append(Ln(ML, 7.02, PAGE_W - MR, 7.02, HAIR, 1.0))
    E.append(T(ML, 7.08, 9.0, 0.35, P(R("让机器狗摔倒后会反思 · 组会汇报", 10.5, INK_FAINT, "sans"), sa=0)))
    E.append(T(PAGE_W - MR - 2.0, 7.08, 2.0, 0.35,
               P(R("%d / 6" % page, 10.5, INK_FAINT, "sans"), align="r", sa=0)))
    return E


# ----------------------------------------------------------------------------- slide 1 cover
def slide_cover():
    E = []
    E += spark(11.55, 2.45, 1.2, SOFT1, 8, 3.0)
    E += spark(11.55, 2.45, 0.66, SOFT2, 8, 3.0)
    E += spark(1.18, 1.63, 0.13, CORAL, 8, 1.9)
    E.append(T(1.42, 1.46, 9, 0.4, P(R("组会汇报 · 2026 年 6 月", 13.5, CORAL, "sans", True), sa=0)))
    E.append(T(1.1, 2.02, 11.2, 1.0, P(R("让机器狗 “摔倒后会反思”", 46, INK, "serif", True), ls=1.0, sa=0)))
    E.append(T(1.1, 3.18, 11.0, 0.9,
               P(R("一个会判断“自己为什么走不动”、并安全选对恢复办法的四足机器人系统", 21, INK_SOFT, "sans"),
                 ls=1.25, sa=0)))
    E.append(Ln(1.13, 4.38, 4.0, 4.38, CORAL, 2.6))
    E.append(T(1.1, 4.58, 11.2, 0.5,
               P(R("项目代号 Kino-VLA　|　宇树 Go2 机器狗　|　高精度物理仿真　|　会看图的 AI 大模型 ＋ 数学安全护盾",
                   13.5, INK_SOFT, "sans"), sa=0)))
    E.append(Ln(1.1, 6.2, PAGE_W - 1.1, 6.2, HAIR, 1.0))
    E.append(T(1.1, 6.36, 7, 0.4, P(R("汇报人：________", 13, INK_FAINT, "sans"), sa=0)))
    E.append(T(PAGE_W - 1.1 - 6.2, 6.36, 6.2, 0.4,
               P(R("已形成 8 页论文初稿（ICRA 会议格式）", 13, INK_FAINT, "sans"), align="r", sa=0)))
    return {"bg": PAGE, "elements": E}


# ----------------------------------------------------------------------------- slide 2 problem
def slide2():
    E = header(2, "为什么需要“会反思”", "机器狗已经很能“扛”，但有一类失败“扛”不过去")
    lx, lw = ML, 6.35

    def point(num, y, title, body):
        out = [T(lx, y, lw, 0.5,
                 P([R(num + "　", 16, CORAL, "sans", True), R(title, 15, INK, "sans", True)], sa=2))]
        out.append(T(lx, y + 0.42, lw, 1.0, P(R(body, 13.5, INK_SOFT, "sans"), ls=1.22, sa=0)))
        return out

    E += point("①", 2.0, "现在的机器狗已经很强",
               "靠强化学习，它能边走边自动适应——打滑、负重、被推，多半能反射式地硬扛过去。这部分做得很好，我们不去改它。")
    E += point("②", 3.2, "但有一类失败，“加力”只会更糟",
               "例：腿被有弹性的绳子缠住，越用力、绳子越像被拉满的弓。正确做法是“后退松劲”，而不是加大力气——需要的是“换个策略”。")
    E += point("③", 4.55, "最棘手：手感一样，却要反着来",
               "有些失败“脚下手感几乎一模一样，却要用相反的办法”。光凭手感分不出来，必须“看一眼”才能判断（见右图）。")

    rx, rw = 7.5, 4.98
    E.append(T(rx, 2.0, rw, 0.4, P(R("一个真实例子：手感相同，处置相反", 13.5, INK, "sans", True), sa=0)))
    # card A
    E.append(Rt(rx, 2.5, rw, 1.12, fill=TINT, line=HAIR, lw=1.0, round=0.08))
    E.append(Rt(rx + rw - 0.6, 2.62, 0.4, 0.4, fill=YELLOW, line=HAIR, lw=0.8, round=0.04))
    E.append(T(rx + 0.22, 2.64, rw - 1.0, 0.4, P(R("脚被胶面粘住（黄色）", 13, CORAL_DEEP, "sans", True), sa=0)))
    E.append(T(rx + 0.22, 3.06, rw - 0.7, 0.5, P(R("正确恢复：后退脱困（越用力越糟）", 12.5, INK, "sans"), sa=0)))
    # card B
    E.append(Rt(rx, 3.72, rw, 1.12, fill=CARD, line=HAIR, lw=1.0, round=0.08))
    E.append(Rt(rx + rw - 0.6, 3.84, 0.4, 0.4, fill=BROWN, line=HAIR, lw=0.8, round=0.04))
    E.append(T(rx + 0.22, 3.86, rw - 1.0, 0.4, P(R("脚陷进烂泥（褐色）", 13, CORAL_DEEP, "sans", True), sa=0)))
    E.append(T(rx + 0.22, 4.28, rw - 0.7, 0.5, P(R("正确恢复：换个步态继续前进", 12.5, INK, "sans"), sa=0)))
    # mini band
    E.append(Rt(rx, 5.0, rw, 0.78, fill=BAND, round=0.06))
    E.append(T(rx + 0.18, 5.0, rw - 0.36, 0.78,
               P(R("关节“手感”几乎相同，颜色完全不同 → 要靠视觉分辨，并记到地图上。", 12.5, INK, "sans", True),
                 align="c", ls=1.18, sa=0), valign="m"))
    # bottom full band
    E.append(Rt(ML, 6.14, CW, 0.66, fill=TINT, round=0.06))
    E.append(T(ML + 0.2, 6.14, CW - 0.4, 0.66,
               P([R("目标：", 14.5, CORAL_DEEP, "sans", True),
                  R("让机器狗像人一样——先判断“是什么情况”，再“对症下药”，而且全程不摔倒。", 14.5, INK, "sans")],
                 align="c", sa=0), valign="m"))
    return {"bg": PAGE, "elements": E}


# ----------------------------------------------------------------------------- slide 3 method
def slide3():
    E = header(3, "方法", "我们的系统：边走边纠错的“四步闭环”")
    bw, gap, by, bh = 2.49, 0.55, 2.0, 2.28
    xs = [ML + i * (bw + gap) for i in range(4)]
    steps = [
        ("①", "实时自检", "每秒上千次", "监测打滑、吃力、走不动；一发现异常，立刻“条件反射”稳住身体（毫秒级），先保证不摔倒。"),
        ("②", "物理状态感知", "看不见也能估", "从关节的“手感”反推地面多滑、负载多重等看不见的物理量，翻译成 AI 能读懂的信息。"),
        ("③", "看图 ＋ 推理", "会看图的 AI", "结合画面、地图和物理手感，判断“到底是什么失败”，写出恢复动作和理由。"),
        ("④", "安全护盾", "数学保证", "AI 给的动作要先过一道安全层，实时把危险指令改安全（不到 1 毫秒），可证明不摔。"),
    ]
    for x, (num, title, sub, body) in zip(xs, steps):
        E.append(Rt(x, by, bw, bh, fill=CARD, line=HAIR, lw=1.0, round=0.08))
        E.append(T(x + 0.18, by + 0.18, bw - 0.36, 0.5,
                   P([R(num + " ", 18, CORAL, "serif", True), R(title, 15, INK, "sans", True)], sa=0)))
        E.append(T(x + 0.18, by + 0.72, bw - 0.36, 0.32, P(R(sub, 11, CORAL_DEEP, "sans", True), sa=0)))
        E.append(Ln(x + 0.18, by + 1.02, x + bw - 0.18, by + 1.02, HAIR, 1.0))
        E.append(T(x + 0.18, by + 1.14, bw - 0.36, bh - 1.26,
                   P(R(body, 12.3, INK_SOFT, "sans"), ls=1.18, sa=0)))
    for x in xs[:3]:
        E.append(T(x + bw, by + bh / 2 - 0.28, gap, 0.56,
                   P(R("→", 20, CORAL, "sans", True), align="c", sa=0), valign="m"))
    # bidirectional link: the planner (box 3) reads & writes the map module below
    cx3 = xs[2] + bw / 2
    E.append(T(cx3 - 1.2, 4.3, 2.4, 0.32,
               P(R("↕  规划器 读图 · 写图", 11, CORAL_DEEP, "sans", True), align="c", sa=0)))
    # semantic / cost map module — the vision layer (spec §7, M5), highlighted
    my = 4.66
    E.append(Rt(ML, my, CW, 1.05, fill=CARD, line=CORAL, lw=1.4, round=0.08))
    E.append(T(ML + 0.22, my + 0.12, CW - 0.44, 0.34,
               P(R("视觉在这里起作用 —— 语义地图（代价地图）", 13.5, CORAL_DEEP, "sans", True), sa=0)))
    E.append(T(ML + 0.22, my + 0.47, CW - 0.44, 0.55,
               P(R("相机把场景分成区域、标出危险，记到一张三维地图上：换个视角也记得，还会推广到长得像的地方（踩穿一块薄冰 → 整片标危险）。规划器读这张图来判断和绕行；当身体真摔了、与画面矛盾时，以身体经历为准、改写地图上的视觉判断。",
                   12.5, INK, "sans"), ls=1.2, sa=0)))
    # dual-rate band
    E.append(Rt(ML, 5.84, CW, 0.64, fill=TINT, round=0.06))
    E.append(T(ML + 0.2, 5.84, CW - 0.4, 0.64,
               P([R("一快一慢配合：", 13.5, CORAL_DEEP, "sans", True),
                  R("慢的“思考”（约 1 秒）想对策；快的“反射”（毫秒级）在思考时不让它倒下——边走边判断边纠正，就是“闭环”。",
                    13.5, INK, "sans")], align="c", ls=1.12, sa=0), valign="m"))
    E.append(T(ML, 6.56, CW, 0.36,
               P(R("注：整套系统都在高精度物理仿真中验证；推理用一个开源、会看图的 AI 大模型（约 40 亿参数，轻量微调）。",
                   11, INK_FAINT, "sans"), align="c", sa=0)))
    return {"bg": PAGE, "elements": E}


# ----------------------------------------------------------------------------- slide 4 data
def slide4():
    E = header(4, "数据与训练", "难点：AI 会“一本正经地胡说”；怎么让它说真话")
    E.append(Rt(ML, 2.0, CW, 0.9, fill=BAND, round=0.06))
    E.append(T(ML + 0.22, 2.0, CW - 0.44, 0.9,
               P([R("问题：", 14.5, CORAL_DEEP, "sans", True),
                  R("让 AI 判断“为什么失败”，它常常自信地猜错（例如一看到黄色就乱归因）。若直接拿它的话去训练，错误也会被一起学进去。",
                    14, INK, "sans")], ls=1.18, sa=0), valign="m"))
    E.append(T(ML, 3.05, CW, 0.4,
               P(R("我们的办法：先在仿真里真实地失败，再用“真值”当判卷老师", 15, INK, "sans", True), sa=0)))
    bw, gap, by, bh = 2.49, 0.55, 3.5, 1.55
    xs = [ML + i * (bw + gap) for i in range(4)]
    steps = [
        ("①", "在仿真里让机器狗真实地“失败”。"),
        ("②", "仿真有“上帝视角”，知道真正的原因和物理参数。"),
        ("③", "调用真实的前沿大模型，写出诊断和恢复理由。"),
        ("④", "用真值判卷：对不上真相、或办法行不通的，一律丢掉。"),
    ]
    for x, (num, body) in zip(xs, steps):
        E.append(Rt(x, by, bw, bh, fill=CARD, line=HAIR, lw=1.0, round=0.08))
        E.append(T(x + 0.18, by + 0.16, bw - 0.36, 0.4, P(R(num, 18, CORAL, "serif", True), sa=0)))
        E.append(T(x + 0.18, by + 0.6, bw - 0.36, bh - 0.72,
                   P(R(body, 12.5, INK, "sans"), ls=1.2, sa=0)))
    for x in xs[:3]:
        E.append(T(x + bw, by + bh / 2 - 0.28, gap, 0.56,
                   P(R("→", 20, CORAL, "sans", True), align="c", sa=0), valign="m"))
    cy, ch, cwl = 5.25, 1.5, 5.6
    E.append(Rt(ML, cy, cwl, ch, fill=TINT, line=HAIR, lw=1.0, round=0.08))
    E.append(T(ML + 0.22, cy + 0.16, cwl - 0.44, 0.4, P(R("数据成果", 14, CORAL_DEEP, "sans", True), sa=0)))
    E.append(T(ML + 0.22, cy + 0.56, cwl - 0.44, ch - 0.64,
               P(R("已收集 413 条真实机器狗失败样本；“真话过滤”丢掉约 21% 的胡说；覆盖多组“看起来一样、其实不同”的成对难题。",
                   13, INK, "sans"), ls=1.22, sa=0)))
    rx2 = ML + cwl + 0.43
    rw2 = CW - cwl - 0.43
    E.append(Rt(rx2, cy, rw2, ch, fill=CARD, line=HAIR, lw=1.0, round=0.08))
    E.append(T(rx2 + 0.22, cy + 0.16, rw2 - 0.44, 0.4, P(R("为什么这一步是“贡献”", 14, CORAL_DEEP, "sans", True), sa=0)))
    E.append(T(rx2 + 0.22, cy + 0.56, rw2 - 0.44, ch - 0.64,
               P(R("这套“真话过滤器”能挡住 AI 的“幻觉”——物理上讲不通的诊断会被自动筛掉。我们已用真实大模型的真实错误验证过它。",
                   13, INK, "sans"), ls=1.22, sa=0)))
    return {"bg": PAGE, "elements": E}


# ----------------------------------------------------------------------------- slide 5 results
def slide5():
    E = header(5, "实验结果", "结果：“会反思”带来不可替代的提升")
    E.append(T(ML, 2.0, 5.6, 0.7,
               P(R("在“看起来一样、其实不同”的难题上，判断“为什么失败”的准确率：", 13.5, INK, "sans", True),
                 ls=1.18, sa=0)))
    base, plot_h, cx0 = 5.18, 2.05, ML + 0.35
    E.append(Ln(cx0 - 0.1, base, cx0 + 4.75, base, HAIR, 1.3))
    bw = 1.25
    b1x, v1 = cx0 + 0.55, 0.974
    b2x, v2 = cx0 + 2.85, 0.237
    h1, h2 = plot_h * v1, plot_h * v2
    E.append(Rt(b1x, base - h1, bw, h1, fill=CORAL))
    E.append(Rt(b2x, base - h2, bw, h2, fill=GRAYBAR))
    E.append(T(b1x - 0.25, base - h1 - 0.62, bw + 0.5, 0.6, P(R("97.4%", 25, CORAL_DEEP, "serif", True), align="c", sa=0), valign="b"))
    E.append(T(b2x - 0.25, base - h2 - 0.62, bw + 0.5, 0.6, P(R("23.7%", 25, INK_SOFT, "serif", True), align="c", sa=0), valign="b"))
    E.append(T(b1x - 0.4, base + 0.1, bw + 0.8, 0.4, P(R("会反思的 AI", 13, INK, "sans", True), align="c", sa=0)))
    E.append(T(b2x - 0.4, base + 0.1, bw + 0.8, 0.4, P(R("传统规则方法", 13, INK_SOFT, "sans", True), align="c", sa=0)))
    E.append(T(ML, 5.78, 5.7, 0.9,
               P([R("高出约 74 个百分点。", 13.5, CORAL_DEEP, "sans", True),
                  R("这类失败靠“手感”或写死的规则都不行，必须“看图 ＋ 推理”。", 13.5, INK, "sans")],
                 ls=1.22, sa=0)))
    rx = 6.75
    rw = CW - (rx - ML)
    cards = [
        ("语义地图（代价地图）", "把危险区域记在三维地图上、换视角也记得；踩穿一块薄冰，就把整片相似区域标危险（必要时用身体经历改写视觉）。"),
        ("安全护盾", "每次决策不到 0.2 毫秒；把危险的速度指令从 2.5 砍到安全范围，全程没有摔倒（且有数学证明）。"),
        ("物理状态感知", "地面摩擦、负载等都估得准（都在容忍误差内）；一旦判出“很滑”，安全余量自动收紧约 77%。"),
        ("全链路打通", "完整跑通“摔倒 → 看＋感 → 诊断 → 安全恢复”；AI 给的恢复动作 100% 能被正确执行。"),
    ]
    cy, chh, cgap = 2.0, 1.12, 0.13
    for title, body in cards:
        E.append(Rt(rx, cy, rw, chh, fill=CARD, line=HAIR, lw=1.0, round=0.07))
        E.append(Rt(rx + 0.2, cy + 0.22, 0.13, 0.13, fill=CORAL))
        E.append(T(rx + 0.46, cy + 0.15, rw - 0.6, 0.32, P(R(title, 13.5, CORAL_DEEP, "sans", True), sa=0)))
        E.append(T(rx + 0.46, cy + 0.49, rw - 0.66, chh - 0.55, P(R(body, 12, INK, "sans"), ls=1.16, sa=0)))
        cy += chh + cgap
    return {"bg": PAGE, "elements": E}


# ----------------------------------------------------------------------------- slide 6 status
def slide6():
    E = header(6, "进展与展望", "进展、诚实的局限与下一步")
    lx, lw = ML, 5.6
    E.append(T(lx, 2.0, lw, 0.4, P(R("进展", 15, CORAL_DEEP, "sans", True), sa=0)))
    nseg = 8
    segw = (lw - 7 * 0.06) / 8
    sy, sh = 2.45, 0.34
    for i in range(nseg):
        x = lx + i * (segw + 0.06)
        if i < 7:
            E.append(Rt(x, sy, segw, sh, fill=CORAL, round=0.03))
        else:
            E.append(Rt(x, sy, segw, sh, fill=None, line=CORAL, lw=1.3, round=0.03))
    E.append(T(lx, 2.9, lw, 0.4, P(R("8 个阶段已完成 7 个；论文（ICRA 会议格式，8 页）初稿已成型。", 12.5, INK, "sans", True), sa=0)))
    E.append(T(lx, 3.28, lw, 0.6,
               P(R("仿真搭建 · 失败情形库 · 监测＋反射 · 安全护盾 · 物理状态感知 · 语义地图 · 数据＋训练 ✓　｜　完整评测（进行中）",
                   11.5, INK_FAINT, "sans"), ls=1.28, sa=0)))
    E.append(T(lx, 3.98, lw, 0.4, P(R("诚实的局限", 15, CORAL_DEEP, "sans", True), sa=0)))
    lims = [
        "数据集还偏小（几百条，属“打通流程”级，目标上千）。",
        "训练与实际部署的“手感”有差异，闭环成功率仍有提升（3 个测试场景成功 1 个）。",
        "“偏好优化”暂未超过基础训练——因为基础训练本身已接近上限。",
    ]
    yy = 4.42
    for t in lims:
        E.append(Rt(lx + 0.02, yy + 0.08, 0.1, 0.1, fill=CORAL))
        E.append(T(lx + 0.28, yy, lw - 0.28, 0.6, P(R(t, 12.5, INK_SOFT, "sans"), ls=1.18, sa=0)))
        yy += 0.56
    rx = 6.75
    rw = CW - (rx - ML)
    E.append(T(rx, 2.0, rw, 0.4, P(R("下一步", 15, CORAL_DEEP, "sans", True), sa=0)))
    nexts = [
        "扩大数据规模（更多仿真场景与样本）。",
        "完善“偏好优化”，争取再上一个台阶。",
        "做完整评测 ＋ 公平的基线对比。",
        "目标投稿 RSS / CoRL / ICRA / IROS 等机器人会议。",
    ]
    yy = 2.5
    for i, t in enumerate(nexts):
        E.append(T(rx, yy, rw, 0.55,
                   P([R(CIRC[i] + " ", 14, CORAL, "sans", True), R(t, 13, INK, "sans")], ls=1.18, sa=0)))
        yy += 0.58
    # right reinforcing card
    E.append(Rt(rx, 5.02, rw, 0.96, fill=BAND, round=0.07))
    E.append(T(rx + 0.22, 5.02, rw - 0.44, 0.96,
               P([R("核心结论：", 13.5, CORAL_DEEP, "sans", True),
                  R("当“手感”分不出原因时，准确率 97.4%（会反思）对 23.7%（规则）。", 13.5, INK, "sans")],
                 ls=1.2, sa=0), valign="m"))
    # closing band
    E.append(Rt(ML, 6.12, CW, 0.74, fill=TINT, round=0.06))
    E.append(T(ML + 0.25, 6.12, CW - 0.5, 0.74,
               P([R("一句话总结：", 14.5, CORAL_DEEP, "sans", True),
                  R("当“手感”无法分辨失败原因时，“看图 ＋ 推理”不可替代；再加上数学安全保证，机器狗就能“一边反思、一边不摔”。",
                    14.5, INK, "sans")], align="c", ls=1.2, sa=0), valign="m"))
    return {"bg": PAGE, "elements": E}


# ----------------------------------------------------------------------------- pptx renderer
def _noshadow(sp):
    try:
        sp.shadow.inherit = False
    except Exception:
        pass


def _setfont(run, font, _bold):
    fam = FAMILY[font]
    run.font.name = fam
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", fam)


def _add_rect(sl, e):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if e["round"] > 0 else MSO_SHAPE.RECTANGLE
    sp = sl.shapes.add_shape(shp, Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    if e["round"] > 0:
        try:
            sp.adjustments[0] = min(0.5, e["round"] / min(e["w"], e["h"]))
        except Exception:
            pass
    if e["fill"]:
        sp.fill.solid()
        sp.fill.fore_color.rgb = hx(e["fill"])
    else:
        sp.fill.background()
    if e["line"]:
        sp.line.color.rgb = hx(e["line"])
        sp.line.width = Pt(e["lw"])
    else:
        sp.line.fill.background()
    _noshadow(sp)


def _add_line(sl, e):
    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                 Inches(e["x1"]), Inches(e["y1"]), Inches(e["x2"]), Inches(e["y2"]))
    cn.line.color.rgb = hx(e["color"])
    cn.line.width = Pt(e["w"])
    _noshadow(cn)


def _add_text(sl, e):
    tb = sl.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[e["valign"]]
    for i, pa in enumerate(e["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[pa["align"]]
        p.line_spacing = pa["ls"]
        p.space_after = Pt(pa["sa"])
        p.space_before = Pt(pa["sb"])
        for rn in pa["runs"]:
            r = p.add_run()
            r.text = rn["t"]
            r.font.size = Pt(rn["size"])
            r.font.bold = rn["bold"]
            r.font.italic = rn["italic"]
            r.font.color.rgb = hx(rn["color"])
            _setfont(r, rn["font"], rn["bold"])


def render_pptx(slides, path):
    prs = Presentation()
    prs.slide_width = Inches(PAGE_W)
    prs.slide_height = Inches(PAGE_H)
    blank = prs.slide_layouts[6]
    for s in slides:
        sl = prs.slides.add_slide(blank)
        bg = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = hx(s["bg"])
        bg.line.fill.background()
        _noshadow(bg)
        for e in s["elements"]:
            if e["type"] == "rect":
                _add_rect(sl, e)
            elif e["type"] == "line":
                _add_line(sl, e)
        for e in s["elements"]:
            if e["type"] == "text":
                _add_text(sl, e)
    prs.save(path)


# ----------------------------------------------------------------------------- pillow preview
_fcache = {}


def pil_font(font, bold, size_pt):
    key = (font, bold, round(size_pt, 1))
    if key not in _fcache:
        pth, idx = FONT_FILES[(font, bold)]
        _fcache[key] = ImageFont.truetype(pth, int(round(size_pt * DPI / 72.0)), index=idx)
    return _fcache[key]


def IN(v):
    return int(round(v * DPI))


def _p_rect(d, e):
    x0, y0 = IN(e["x"]), IN(e["y"])
    x1, y1 = IN(e["x"] + e["w"]), IN(e["y"] + e["h"])
    fill = rgb(e["fill"]) if e["fill"] else None
    line = rgb(e["line"]) if e["line"] else None
    w = max(1, int(round(e["lw"] * DPI / 72.0)))
    if e["round"] > 0:
        d.rounded_rectangle([x0, y0, x1, y1], radius=max(1, IN(e["round"])), fill=fill, outline=line, width=w)
    else:
        d.rectangle([x0, y0, x1, y1], fill=fill, outline=line, width=w)


def _p_line(d, e):
    d.line([(IN(e["x1"]), IN(e["y1"])), (IN(e["x2"]), IN(e["y2"]))],
           fill=rgb(e["color"]), width=max(1, int(round(e["w"] * DPI / 72.0))))


def _p_text(d, e):
    pad = IN(0.1)
    top_pad = IN(0.04)
    box_x = IN(e["x"]) + pad
    box_y = IN(e["y"]) + top_pad
    box_w = IN(e["w"]) - 2 * pad
    box_h = IN(e["h"])
    para_lines = []
    total_h = 0.0
    for pa in e["paras"]:
        atoms = []
        for rn in pa["runs"]:
            f = pil_font(rn["font"], rn["bold"], rn["size"])
            for ch in rn["t"]:
                atoms.append((ch, f, rn["size"], rgb(rn["color"])))
        lines = []
        cur = []
        cur_w = 0.0
        for a in atoms:
            ch, f, _sz, _c = a
            if ch == "\n":
                lines.append(cur)
                cur = []
                cur_w = 0.0
                continue
            cw = d.textlength(ch, font=f)
            if cur and cur_w + cw > box_w:
                lines.append(cur)
                cur = [a]
                cur_w = cw
            else:
                cur.append(a)
                cur_w += cw
        if cur:
            lines.append(cur)
        if not lines:
            lines = [[]]
        para_lines.append((lines, pa["ls"], pa["align"], pa["sa"]))
        for ln in lines:
            mh = max([sz for (_, _, sz, _) in ln], default=pa["runs"][0]["size"])
            total_h += mh * DPI / 72.0 * pa["ls"]
        total_h += pa["sa"] * DPI / 72.0
    if e["valign"] == "m":
        oy = box_y + max(0, (box_h - 2 * top_pad - total_h) / 2)
    elif e["valign"] == "b":
        oy = box_y + max(0, (box_h - 2 * top_pad - total_h))
    else:
        oy = box_y
    y = oy
    for (lines, ls, align, sa) in para_lines:
        for ln in lines:
            line_w = sum(d.textlength(ch, font=f) for (ch, f, _, _) in ln)
            mh = max([sz for (_, _, sz, _) in ln], default=12)
            lh = mh * DPI / 72.0 * ls
            if align == "c":
                x = box_x + max(0, (box_w - line_w) / 2)
            elif align == "r":
                x = box_x + max(0, (box_w - line_w))
            else:
                x = box_x
            for (ch, f, _sz, color) in ln:
                d.text((x, y), ch, font=f, fill=color, anchor="la")
                x += d.textlength(ch, font=f)
            y += lh
        y += sa * DPI / 72.0


def render_preview(slides, outdir):
    os.makedirs(outdir, exist_ok=True)
    paths = []
    for n, s in enumerate(slides, 1):
        img = Image.new("RGB", (IN(PAGE_W), IN(PAGE_H)), rgb(s["bg"]))
        d = ImageDraw.Draw(img)
        for e in s["elements"]:
            if e["type"] == "rect":
                _p_rect(d, e)
            elif e["type"] == "line":
                _p_line(d, e)
        for e in s["elements"]:
            if e["type"] == "text":
                _p_text(d, e)
        p = os.path.join(outdir, "slide_%d.png" % n)
        img.save(p)
        paths.append(p)
    return paths


# ----------------------------------------------------------------------------- main
def main():
    here = os.path.dirname(os.path.abspath(__file__))
    slides = [slide_cover(), slide2(), slide3(), slide4(), slide5(), slide6()]
    pptx_path = os.path.join(here, "KinoVLA_组会汇报.pptx")
    render_pptx(slides, pptx_path)
    preview_dir = os.path.join(here, "preview")
    paths = render_preview(slides, preview_dir)
    pdf_path = os.path.join(here, "KinoVLA_组会汇报_preview.pdf")
    try:
        Image.init()  # force-register codec plugins (lazy JPEG handler for PDF)
        imgs = [Image.open(p).convert("RGB") for p in paths]
        imgs[0].save(pdf_path, save_all=True, append_images=imgs[1:], resolution=DPI)
        print("PDF  :", pdf_path)
    except Exception as exc:  # PDF preview is a bonus; never block pptx/png
        print("PDF  : (skipped:", exc, ")")
    print("PPTX :", pptx_path)
    for p in paths:
        print("PNG  :", p)


if __name__ == "__main__":
    main()
