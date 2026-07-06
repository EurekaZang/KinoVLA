#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""KiNO paper Fig. 1 — the closed-loop system diagram.

IEEE / top-robotics-venue aesthetic: white background, compact, high density,
color-coded modules, labelled rates, a line-style legend, four edge types
(data flow / fast reflex / privileged distill / mu-hat coupling). English labels.

Single source -> editable .pptx + a publication-grade supersampled PNG (no
LibreOffice on this box, so the PNG is rendered directly from the same spec).
"""
import math
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

# ------------------------------------------------------------------ palette (IEEE, restrained)
WHITE = "FFFFFF"
INK = "13171C"        # near-black title text
SUB = "59616B"        # secondary grey text
ARROW = "2C313A"      # data-flow arrows
DASH = "8A9099"       # privileged distillation (dashed)
RED = "C0392B"        # fast reflex loop
BLUE_F, BLUE_L = "EAF1FB", "2F6CB0"   # online reasoning layers
GRAY_F, GRAY_L = "ECEEF1", "6C737B"   # plant / robot
AMBER_F, AMBER_L = "FBEFDA", "BE7E2C"  # privileged / auxiliary
LEG_LINE = "D8DCE1"

NOTO = "/usr/share/fonts/truetype/noto/NotoSans-%s.ttf"
FONT_FILES_FIG = {"reg": NOTO % "Regular", "bold": NOTO % "Bold",
                  "ital": NOTO % "Italic", "boldital": NOTO % "BoldItalic"}
PPTX_FONT = "Noto Sans"

W, H = 12.0, 4.45          # design canvas (inches); ~2.70:1 -> ~2.65in tall at \textwidth
FINAL_DPI = 300
SS = 2                      # supersample factor
RENDER_DPI = FINAL_DPI * SS


def hx(h):
    return RGBColor.from_string(h)


def rgb(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def IN(v):
    return int(round(v * RENDER_DPI))


# ------------------------------------------------------------------ element builders
def R(t, size, color=INK, style="reg"):
    return dict(t=t, size=size, color=color, style=style)


def P(run, align="c", ls=1.0, sa=0.0):
    return dict(runs=[run], align=align, ls=ls, sa=sa)


def Txt(x, y, w, h, paras, valign="m"):
    if isinstance(paras, dict):
        paras = [paras]
    return dict(type="text", x=x, y=y, w=w, h=h, paras=paras, valign=valign)


def Rect(x, y, w, h, fill=None, line=None, lw=1.4, round=0.06):
    return dict(type="rect", x=x, y=y, w=w, h=h, fill=fill, line=line, lw=lw, round=round)


def Path(pts, color=ARROW, w=1.6, dash=None, head=True, tail=False):
    return dict(type="path", pts=pts, color=color, w=w, dash=dash, head=head, tail=tail)


def module(x, y, w, h, fill, line, title, rate=None, detail=None, sub=None,
           rate_color=None, title_size=15.5):
    """A labelled box: bold title (+ optional rate / detail / italic sub), centred."""
    els = [Rect(x, y, w, h, fill=fill, line=line, lw=1.5)]
    paras = [P(R(title, title_size, INK, "bold"), sa=2.0)]
    if rate:
        paras.append(P(R(rate, 11.5, rate_color or AMBER_L, "boldital"), sa=2.0))
    if detail:
        paras.append(P(R(detail, 11.0, SUB, "reg")))
    if sub:
        paras.append(P(R(sub, 10.5, SUB, "ital")))
    els.append(Txt(x, y, w, h, paras, valign="m"))
    return els


# ------------------------------------------------------------------ layout
def build():
    E = []

    # --- main pipeline row -------------------------------------------------
    bw, bh, my = 1.70, 0.98, 1.62
    xs = [0.30, 2.725, 5.15, 7.575, 10.0]
    cx = [x + bw / 2 for x in xs]
    arr_y = my + bh / 2          # 2.11
    bot = my + bh                # 2.60

    E += module(xs[0], my, bw, bh, GRAY_F, GRAY_L, "Unitree Go2", sub="Isaac Lab")
    E += module(xs[1], my, bw, bh, BLUE_F, BLUE_L, "Kino-Monitor", rate="1 kHz",
                detail="slip · tracking error", rate_color=BLUE_L)
    E += module(xs[2], my, bw, bh, BLUE_F, BLUE_L, "Kino-Tokens", rate="500 ms",
                detail="1D-CNN + Perceiver", rate_color=BLUE_L)
    E += module(xs[3], my, bw, bh, BLUE_F, BLUE_L, "VLA Planner", rate="~1 Hz",
                detail="Qwen3-VL-4B + LoRA", rate_color=BLUE_L)
    E += module(xs[4], my, bw, bh, BLUE_F, BLUE_L, "CBF-QP Shield", rate="< 0.2 ms",
                detail="+ Primitive Compiler", rate_color=BLUE_L)

    # forward data-flow arrows + labels
    labels = ["proprio.", "window", "tokens", "<Action>"]
    for i in range(4):
        x0, x1 = xs[i] + bw, xs[i + 1]
        E.append(Path([(x0, arr_y), (x1, arr_y)], color=ARROW, w=1.7, head=True))
        E.append(Txt((x0 + x1) / 2 - 0.45, 1.80, 0.90, 0.24,
                     P(R(labels[i], 11.0, ARROW, "reg")), valign="b"))

    # --- privileged physics theta (top, above Kino-Tokens) ----------------
    th_w, th_h = 2.5, 0.66
    th_x, th_y = cx[2] - th_w / 2, 0.42
    E += module(th_x, th_y, th_w, th_h, AMBER_F, AMBER_L, "Privileged Physics  θ",
                sub="simulator truth · training-only", title_size=14.0)
    E.append(Path([(cx[2], th_y + th_h), (cx[2], my)], color=DASH, w=1.4, dash="dash", head=True))
    E.append(Txt(cx[2] + 0.08, 1.18, 0.9, 0.22, P(R("distill", 10.5, SUB, "ital"), align="l"), valign="m"))

    # --- mu-hat friction coupling (Kino-Tokens -> Shield, orthogonal arc) --
    cpx0, cpx1, cpy = 6.55, 10.5, 1.32
    E.append(Path([(cpx0, my), (cpx0, cpy), (cpx1, cpy), (cpx1, my)],
                  color=AMBER_L, w=1.5, dash="dot", head=True))
    E.append(Txt(7.1, 1.06, 2.8, 0.22, P(R("μ̂  friction bound", 11.0, AMBER_L, "bold")), valign="m"))

    # --- reflex loop (below Kino-Monitor), fast red loop ------------------
    rf_w, rf_h = 2.10, 0.74
    rf_x, rf_y = cx[1] - rf_w / 2, 2.98
    E += module(rf_x, rf_y, rf_w, rf_h, AMBER_F, AMBER_L, "Reflex Loop",
                sub="< 20 ms · self-stabilize", title_size=14.0)
    E.append(Path([(cx[1], bot), (cx[1], rf_y)], color=RED, w=1.9, head=True))
    E.append(Txt(cx[1] + 0.10, 2.66, 1.0, 0.22, P(R("anomaly", 10.5, RED, "bold"), align="l"), valign="m"))
    E.append(Path([(rf_x, rf_y + rf_h / 2), (1.70, rf_y + rf_h / 2), (1.70, bot)],
                  color=RED, w=1.9, head=True))

    # --- semantic traversability map (below VLA Planner) ------------------
    mp_w, mp_h = 2.90, 0.74
    mp_x, mp_y = cx[3] - mp_w / 2, 2.98
    E += module(mp_x, mp_y, mp_w, mp_h, AMBER_F, AMBER_L, "Semantic Traversability Map",
                sub="persistent across viewpoints", title_size=13.5)
    E.append(Path([(cx[3], bot), (cx[3], mp_y)], color=ARROW, w=1.6, head=True, tail=True))
    E.append(Txt(cx[3] + 0.10, 2.66, 1.2, 0.22, P(R("read · write", 10.0, SUB, "ital"), align="l"), valign="m"))

    # RGB-D sensor input -> map
    rd_w, rd_h = 1.28, 0.60
    rd_x, rd_y = 4.96, mp_y + mp_h / 2 - rd_h / 2
    E += module(rd_x, rd_y, rd_w, rd_h, GRAY_F, GRAY_L, "RGB-D", sub="30 Hz", title_size=13.0)
    E.append(Path([(rd_x + rd_w, mp_y + mp_h / 2), (mp_x, mp_y + mp_h / 2)], color=ARROW, w=1.5, head=True))

    # --- closed-loop feedback (Shield -> robot, along the bottom) ----------
    fb_y = 4.12
    E.append(Path([(cx[4], bot), (cx[4], fb_y), (1.10, fb_y), (1.10, bot)],
                  color=ARROW, w=1.7, head=True))
    E.append(Txt(2.6, 3.86, 6.6, 0.24,
                 P(R("filtered velocity command   +   execution report", 11.0, ARROW, "reg")), valign="b"))

    # --- legend (top-left) -------------------------------------------------
    lx, ly, lw, lh = 0.30, 0.36, 2.32, 0.92
    E.append(Rect(lx, ly, lw, lh, fill=WHITE, line=LEG_LINE, lw=1.0, round=0.04))
    E.append(Txt(lx + 0.14, ly + 0.06, lw - 0.2, 0.2, P(R("Edge types", 10.5, INK, "bold"), align="l"), valign="t"))
    rows = [("data flow", ARROW, None), ("reflex   < 20 ms", RED, None),
            ("privileged distill", DASH, "dash"), ("μ̂ coupling", AMBER_L, "dot")]
    ry = ly + 0.30
    for txt, col, dash in rows:
        E.append(Path([(lx + 0.16, ry + 0.075), (lx + 0.66, ry + 0.075)], color=col, w=1.7, dash=dash, head=True))
        E.append(Txt(lx + 0.78, ry - 0.02, lw - 0.82, 0.2, P(R(txt, 10.0, INK, "reg"), align="l"), valign="t"))
        ry += 0.15

    return E


# ------------------------------------------------------------------ pillow render
_fc = {}


def pil_font(style, size):
    k = (style, round(size, 1))
    if k not in _fc:
        _fc[k] = ImageFont.truetype(FONT_FILES_FIG[style], int(round(size * RENDER_DPI / 72.0)))
    return _fc[k]


def _dashed(d, p0, p1, col, w, on, off):
    x0, y0 = p0
    x1, y1 = p1
    dist = math.hypot(x1 - x0, y1 - y0) or 1.0
    ux, uy = (x1 - x0) / dist, (y1 - y0) / dist
    pos = 0.0
    while pos < dist:
        seg = min(on, dist - pos)
        a = (x0 + ux * pos, y0 + uy * pos)
        b = (x0 + ux * (pos + seg), y0 + uy * (pos + seg))
        d.line([a, b], fill=col, width=w)
        pos += on + off


def _arrowhead(d, tip, frm, col):
    L, Wd = IN(0.10), IN(0.060)
    dx, dy = tip[0] - frm[0], tip[1] - frm[1]
    n = math.hypot(dx, dy) or 1.0
    ux, uy = dx / n, dy / n
    bx, by = tip[0] - ux * L, tip[1] - uy * L
    px, py = -uy, ux
    d.polygon([tip, (bx + px * Wd, by + py * Wd), (bx - px * Wd, by - py * Wd)], fill=col)


def _p_path(d, e):
    pts = [(IN(x), IN(y)) for (x, y) in e["pts"]]
    col = rgb(e["color"])
    w = max(1, int(round(e["w"] * RENDER_DPI / 72.0)))
    for i in range(len(pts) - 1):
        p0, p1 = pts[i], pts[i + 1]
        if e["dash"] == "dash":
            _dashed(d, p0, p1, col, w, IN(0.075), IN(0.05))
        elif e["dash"] == "dot":
            _dashed(d, p0, p1, col, max(w, IN(0.016)), IN(0.018), IN(0.034))
        else:
            d.line([p0, p1], fill=col, width=w)
    if e["head"]:
        _arrowhead(d, pts[-1], pts[-2], col)
    if e["tail"]:
        _arrowhead(d, pts[0], pts[1], col)


def _p_rect(d, e):
    x0, y0, x1, y1 = IN(e["x"]), IN(e["y"]), IN(e["x"] + e["w"]), IN(e["y"] + e["h"])
    fill = rgb(e["fill"]) if e["fill"] else None
    line = rgb(e["line"]) if e["line"] else None
    w = max(1, int(round(e["lw"] * RENDER_DPI / 72.0)))
    if e["round"] > 0:
        d.rounded_rectangle([x0, y0, x1, y1], radius=max(1, IN(e["round"])), fill=fill, outline=line, width=w)
    else:
        d.rectangle([x0, y0, x1, y1], fill=fill, outline=line, width=w)


def _wrap_words(text, font, max_w, d):
    out = []
    for part in text.split("\n"):
        words = part.split(" ")
        cur = ""
        for wd in words:
            trial = wd if cur == "" else cur + " " + wd
            if cur == "" or d.textlength(trial, font=font) <= max_w:
                cur = trial
            else:
                out.append(cur)
                cur = wd
        out.append(cur)
    return out or [""]


def _p_text(d, e):
    padx, padyt = IN(0.06), IN(0.04)
    bx, bw = IN(e["x"]) + padx, IN(e["w"]) - 2 * padx
    by, bh = IN(e["y"]), IN(e["h"])
    lines = []
    for pa in e["paras"]:
        rn = pa["runs"][0]
        font = pil_font(rn["style"], rn["size"])
        wl = _wrap_words(rn["t"], font, bw, d)
        for i, ln in enumerate(wl):
            extra = pa["sa"] if i == len(wl) - 1 else 0.0
            lines.append((ln, font, rn["size"], rgb(rn["color"]), pa["align"], pa["ls"], extra))
    total = sum(sz * RENDER_DPI / 72.0 * ls + sa * RENDER_DPI / 72.0
                for (_, _, sz, _, _, ls, sa) in lines)
    if e["valign"] == "m":
        y = by + max(0, (bh - total) / 2)
    elif e["valign"] == "b":
        y = by + max(0, bh - total) - padyt
    else:
        y = by + padyt
    for (txt, font, sz, col, align, ls, sa) in lines:
        if align == "c":
            d.text((bx + bw / 2, y), txt, font=font, fill=col, anchor="ma")
        elif align == "r":
            d.text((bx + bw, y), txt, font=font, fill=col, anchor="ra")
        else:
            d.text((bx, y), txt, font=font, fill=col, anchor="la")
        y += sz * RENDER_DPI / 72.0 * ls + sa * RENDER_DPI / 72.0


def render_png(E, path):
    img = Image.new("RGB", (IN(W), IN(H)), rgb(WHITE))
    d = ImageDraw.Draw(img)
    for e in E:
        if e["type"] == "rect":
            _p_rect(d, e)
    for e in E:
        if e["type"] == "path":
            _p_path(d, e)
    for e in E:
        if e["type"] == "text":
            _p_text(d, e)
    img = img.resize((int(W * FINAL_DPI), int(H * FINAL_DPI)), Image.LANCZOS)
    img.save(path, dpi=(FINAL_DPI, FINAL_DPI))
    return img.size


# ------------------------------------------------------------------ pptx render
def _noshadow(sp):
    try:
        sp.shadow.inherit = False
    except Exception:
        pass


def _append(ln, tag, attrs):
    el = ln.makeelement(qn(tag), attrs)
    ln.append(el)
    return el


def _add_rect(sl, e):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if e["round"] > 0 else MSO_SHAPE.RECTANGLE
    sp = sl.shapes.add_shape(shp, Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    if e["round"] > 0:
        try:
            sp.adjustments[0] = min(0.5, e["round"] / min(e["w"], e["h"]))
        except Exception:
            pass
    if e["fill"]:
        sp.fill.solid()
        sp.fill.fore_color.rgb = hx(e["fill"])
    else:
        sp.fill.background()
    if e["line"]:
        sp.line.color.rgb = hx(e["line"])
        sp.line.width = Pt(e["lw"])
    else:
        sp.line.fill.background()
    _noshadow(sp)


def _add_path(sl, e):
    pts = e["pts"]
    n = len(pts)
    for i in range(n - 1):
        p0, p1 = pts[i], pts[i + 1]
        cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                     Inches(p0[0]), Inches(p0[1]), Inches(p1[0]), Inches(p1[1]))
        cn.line.color.rgb = hx(e["color"])
        cn.line.width = Pt(e["w"])
        _noshadow(cn)
        ln = cn.line._get_or_add_ln()
        if e["dash"] == "dash":
            _append(ln, "a:prstDash", {"val": "dash"})
        elif e["dash"] == "dot":
            _append(ln, "a:prstDash", {"val": "sysDot"})
        if e["tail"] and i == 0:
            _append(ln, "a:headEnd", {"type": "triangle", "w": "med", "len": "med"})
        if e["head"] and i == n - 2:
            _append(ln, "a:tailEnd", {"type": "triangle", "w": "med", "len": "med"})


def _add_text(sl, e):
    tb = sl.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.04))
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[e["valign"]]
    for i, pa in enumerate(e["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"c": PP_ALIGN.CENTER, "l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT}[pa["align"]]
        p.line_spacing = pa["ls"]
        p.space_after = Pt(pa["sa"])
        p.space_before = Pt(0)
        rn = pa["runs"][0]
        r = p.add_run()
        r.text = rn["t"]
        r.font.size = Pt(rn["size"])
        r.font.bold = "bold" in rn["style"]
        r.font.italic = "ital" in rn["style"]
        r.font.color.rgb = hx(rn["color"])
        r.font.name = PPTX_FONT


def render_pptx(E, path):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = hx(WHITE)
    bg.line.fill.background()
    _noshadow(bg)
    for e in E:
        if e["type"] == "rect":
            _add_rect(sl, e)
    for e in E:
        if e["type"] == "path":
            _add_path(sl, e)
    for e in E:
        if e["type"] == "text":
            _add_text(sl, e)
    prs.save(path)


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    paper = os.path.abspath(os.path.join(here, ".."))
    E = build()
    pptx_path = os.path.join(here, "fig1_system.pptx")
    png_path = os.path.join(paper, "fig1_system.png")
    render_pptx(E, pptx_path)
    size = render_png(E, png_path)
    print("PPTX :", pptx_path)
    print("PNG  :", png_path, size)


if __name__ == "__main__":
    main()
